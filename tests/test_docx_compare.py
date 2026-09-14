from __future__ import annotations

import json
import os
import subprocess
import sys
from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document as PyDocxDocument
from lxml import etree

from docxtor import (
    RevisionAuthor,
    RevisionRange,
    accept_all_revisions_bytes,
    compare_docx_documents,
)
from docxtor.docx_revision_mutations import replace_revision

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS = {"w": _W}


def _docx(paragraphs: tuple[str, ...] = ()) -> bytes:
    document = PyDocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _rewrite_document(data: bytes, edit) -> bytes:
    with ZipFile(BytesIO(data)) as source:
        entries = [(item.filename, source.read(item.filename)) for item in source.infolist()]
    output = BytesIO()
    with ZipFile(output, mode="w", compression=ZIP_DEFLATED) as target:
        for name, payload in entries:
            if name == "word/document.xml":
                root = etree.fromstring(payload)
                edit(root)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
            target.writestr(name, payload)
    return output.getvalue()


def _docx_with_footnote(note_text: str) -> bytes:
    def add_reference(root) -> None:
        paragraph = root.xpath("//w:p", namespaces=_NS)[0]
        run = etree.SubElement(paragraph, f"{{{_W}}}r")
        etree.SubElement(run, f"{{{_W}}}footnoteReference", {f"{{{_W}}}id": "1"})

    data = _rewrite_document(_docx(("Body paragraph.",)), add_reference)
    with ZipFile(BytesIO(data)) as source:
        entries = {item.filename: source.read(item.filename) for item in source.infolist()}

    relationships_name = "word/_rels/document.xml.rels"
    entries[relationships_name] = entries[relationships_name].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdFootnotes" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        b'Target="footnotes.xml"/></Relationships>',
    )
    entries["[Content_Types].xml"] = entries["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/word/footnotes.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        b"</Types>",
    )
    entries["word/footnotes.xml"] = (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="-1" w:type="separator"><w:p/></w:footnote>'
        '<w:footnote w:id="0" w:type="continuationSeparator"><w:p/></w:footnote>'
        f'<w:footnote w:id="1"><w:p><w:r><w:t>{note_text}</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    ).encode()
    output = BytesIO()
    with ZipFile(output, mode="w", compression=ZIP_DEFLATED) as target:
        for name, payload in entries.items():
            target.writestr(name, payload)
    return output.getvalue()


def test_identical_documents_have_stable_empty_comparison_and_json_data() -> None:
    source = _docx(("First paragraph.", "Second paragraph."))

    result = compare_docx_documents(source, source)
    repeated = compare_docx_documents(source, source)

    assert result.left.sha256 == sha256(source).hexdigest()
    assert result.right.sha256 == sha256(source).hexdigest()
    assert result.comparison_id == repeated.comparison_id
    assert result.text_changes == ()
    assert result.comment_changes == ()
    assert result.revision_events == ()
    assert result.formatting_changes == ()
    assert [block.raw_text for block in result.left.blocks] == [
        "First paragraph.",
        "Second paragraph.",
    ]
    assert [pair.pair_id for pair in result.block_pairs] == [
        pair.pair_id for pair in repeated.block_pairs
    ]
    json.dumps(asdict(result))


def test_plain_text_insertion_and_deletion_have_side_specific_anchors() -> None:
    left = _docx(("Keep this clause.", "Remove this clause.", "Final clause."))
    right = _docx(("Add this clause.", "Keep this clause.", "Final clause."))

    result = compare_docx_documents(left, right)

    inserted = next(change for change in result.text_changes if change.kind.value == "inserted")
    deleted = next(change for change in result.text_changes if change.kind.value == "deleted")
    assert inserted.left is None and inserted.right is not None
    assert inserted.right.text == "Add this clause."
    assert inserted.right.start_offset == 0
    assert inserted.right.block_id in {block.id for block in result.right.blocks}
    assert deleted.left is not None and deleted.right is None
    assert deleted.left.text == "Remove this clause."
    assert deleted.left.start_offset == 0
    assert deleted.left.block_id in {block.id for block in result.left.blocks}


def test_repeated_identical_text_changes_have_unique_stable_ids() -> None:
    left = _docx(("Pay 14 days; deliver 14 days.",))
    right = _docx(("Pay 30 days; deliver 30 days.",))

    result = compare_docx_documents(left, right)
    repeated = compare_docx_documents(left, right)

    assert len(result.text_changes) == 2
    assert [change.left.start_offset for change in result.text_changes if change.left] == [4, 21]
    change_ids = [change.change_id for change in result.text_changes]
    assert len(set(change_ids)) == 2
    assert change_ids == [change.change_id for change in repeated.text_changes]


def test_repeated_paragraphs_remain_aligned_after_one_inserted_paragraph() -> None:
    left = _docx(("Repeated paragraph.", "Repeated paragraph.", "Closing paragraph."))
    right = _docx(
        (
            "Added paragraph.",
            "Repeated paragraph.",
            "Repeated paragraph.",
            "Closing paragraph.",
        )
    )

    result = compare_docx_documents(left, right)
    pairs_by_id = {pair.pair_id: pair for pair in result.block_pairs}
    repeated_pairs = [
        pair
        for pair in result.block_pairs
        if pair.left_block_id
        and pair.right_block_id
        and next(block for block in result.left.blocks if block.id == pair.left_block_id).text
        == "Repeated paragraph."
    ]

    assert len(repeated_pairs) == 2
    assert all(pair.left_block_id and pair.right_block_id for pair in repeated_pairs)
    assert len(pairs_by_id) == 4
    assert any(change.kind.value == "inserted" for change in result.text_changes)
    assert not any(change.kind.value == "deleted" for change in result.text_changes)


def test_accepting_tracked_replacement_keeps_text_diff_empty_and_reports_lifecycle() -> None:
    original = _docx(("The notice period is 14 days.",))
    reviewed = replace_revision(
        original,
        RevisionRange("body:p:0", 21, 23, expected_text="14"),
        "30",
        RevisionAuthor("Word reviewer"),
    )[1].data
    accepted = accept_all_revisions_bytes(reviewed, drop_comments=False).output_bytes

    result = compare_docx_documents(reviewed, accepted)

    assert result.left.blocks[0].text == result.right.blocks[0].text
    assert result.text_changes == ()
    assert result.formatting_changes == ()
    assert {event.status.value for event in result.revision_events} == {"resolved"}
    assert {event.resolution_evidence.value for event in result.revision_events} == {
        "accepted_text_equal"
    }
    assert {span.revision_kind for span in result.left.blocks[0].spans if span.revision_id} == {
        "del",
        "ins",
    }


def test_comment_only_change_is_separate_and_keeps_table_anchor_coordinates() -> None:
    plain = PyDocxDocument()
    plain.add_paragraph("Body paragraph.")
    cell = plain.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Table cell text.")
    plain_stream = BytesIO()
    plain.save(plain_stream)

    commented = PyDocxDocument()
    commented.add_paragraph("Body paragraph.")
    cell = commented.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Table cell text.")
    commented.add_comment(
        runs=cell.paragraphs[0].runs[0],
        text="Please verify this table value.",
        author="Reviewer",
    )
    commented_stream = BytesIO()
    commented.save(commented_stream)

    result = compare_docx_documents(plain_stream.getvalue(), commented_stream.getvalue())

    assert result.text_changes == ()
    assert len(result.comment_changes) == 1
    assert result.comment_changes[0].status.value == "added"
    comment = result.right.comments[0]
    assert comment.anchor is not None
    assert comment.anchor.locator.startswith("table:0:r:0:c:0:")
    block = next(block for block in result.right.blocks if block.id == comment.anchor.block_id)
    assert block.coordinate.table_index == 0
    assert block.coordinate.row_index == 0
    assert block.coordinate.cell_index == 0


def test_formatting_only_change_is_reported_separately_from_text() -> None:
    left = PyDocxDocument()
    left.add_paragraph("Same visible text.")
    left_stream = BytesIO()
    left.save(left_stream)

    right = PyDocxDocument()
    paragraph = right.add_paragraph()
    paragraph.add_run("Same visible text.").bold = True
    right_stream = BytesIO()
    right.save(right_stream)

    result = compare_docx_documents(left_stream.getvalue(), right_stream.getvalue())

    assert result.text_changes == ()
    assert len(result.formatting_changes) == 1
    assert result.formatting_changes[0].left_block_id is not None
    assert result.formatting_changes[0].right_block_id is not None


def test_formatting_change_on_shared_characters_with_text_edit_is_diagnosed() -> None:
    left_document = PyDocxDocument()
    paragraph = left_document.add_paragraph()
    paragraph.add_run("ab").bold = True
    paragraph.add_run("cd")
    left_stream = BytesIO()
    left_document.save(left_stream)

    right_document = PyDocxDocument()
    paragraph = right_document.add_paragraph()
    paragraph.add_run("abc").bold = True
    paragraph.add_run("de")
    right_stream = BytesIO()
    right_document.save(right_stream)

    result = compare_docx_documents(left_stream.getvalue(), right_stream.getvalue())

    assert result.text_changes
    assert result.formatting_changes == ()
    assert result.coverage.value == "incomplete"
    assert any(
        diagnostic.code == "run_formatting_not_isolated_with_text_change"
        for diagnostic in result.diagnostics
    )


def test_text_length_change_with_same_run_style_is_not_a_formatting_change() -> None:
    left = _docx(("Short clause.",))
    right = _docx(("A much longer clause.",))

    result = compare_docx_documents(left, right)

    assert result.text_changes
    assert result.formatting_changes == ()


def test_changed_run_formatting_boundary_is_not_lost_when_text_is_unchanged() -> None:
    left = PyDocxDocument()
    paragraph = left.add_paragraph()
    paragraph.add_run("ab").bold = True
    paragraph.add_run("cd").bold = False
    left_stream = BytesIO()
    left.save(left_stream)

    right = PyDocxDocument()
    paragraph = right.add_paragraph()
    paragraph.add_run("abc").bold = True
    paragraph.add_run("d").bold = False
    right_stream = BytesIO()
    right.save(right_stream)

    result = compare_docx_documents(left_stream.getvalue(), right_stream.getvalue())

    assert result.text_changes == ()
    assert any(change.scope == "run_properties" for change in result.formatting_changes)


def test_body_table_body_blocks_keep_source_order() -> None:
    document = PyDocxDocument()
    document.add_paragraph("Before table.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Inside table."
    document.add_paragraph("After table.")
    stream = BytesIO()
    document.save(stream)

    result = compare_docx_documents(stream.getvalue(), stream.getvalue())

    assert [block.coordinate.container_id for block in result.left.blocks] == [
        "body:p:0",
        "table:0:r:0:c:0:p:0",
        "body:p:1",
    ]
    assert [block.order for block in result.left.blocks] == [0, 1, 2]


def test_middle_insertion_keeps_navigation_pairs_in_document_order() -> None:
    left = _docx(("First.", "Second.", "Third."))
    right = _docx(("First.", "Inserted in the middle.", "Second.", "Third."))

    result = compare_docx_documents(left, right)

    assert [pair.right_block_id for pair in result.block_pairs] == [
        block.id for block in result.right.blocks
    ]
    assert [block.pair_id for block in result.right.blocks] == [
        pair.pair_id for pair in result.block_pairs
    ]
    assert [change.right.text for change in result.text_changes] == ["Inserted in the middle."]


def test_shifted_table_ordinal_emits_incomplete_alignment_diagnostic() -> None:
    before_document = PyDocxDocument()
    before_document.add_table(rows=1, cols=1).cell(0, 0).text = "Original table."
    before_stream = BytesIO()
    before_document.save(before_stream)

    after_document = PyDocxDocument()
    after_document.add_table(rows=1, cols=1).cell(0, 0).text = "New table."
    after_document.add_table(rows=1, cols=1).cell(0, 0).text = "Original table."
    after_stream = BytesIO()
    after_document.save(after_stream)

    result = compare_docx_documents(before_stream.getvalue(), after_stream.getvalue())

    assert result.coverage.value == "incomplete"
    assert any(
        diagnostic.code == "table_ordinal_alignment_coarse" for diagnostic in result.diagnostics
    )


def test_identical_repeated_tables_do_not_report_shifted_ordinals() -> None:
    document = PyDocxDocument()
    for _ in range(2):
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Repeated table."
    stream = BytesIO()
    document.save(stream)

    result = compare_docx_documents(stream.getvalue(), stream.getvalue())

    assert result.coverage.value == "complete"
    assert not any(
        diagnostic.code == "table_ordinal_alignment_coarse" for diagnostic in result.diagnostics
    )


def test_story_pair_ids_are_stable_across_calls_and_subprocess(tmp_path) -> None:
    document = PyDocxDocument()
    document.add_paragraph("Body paragraph.")
    document.sections[0].header.paragraphs[0].text = "Header paragraph."
    document.sections[0].footer.paragraphs[0].text = "Footer paragraph."
    source = tmp_path / "story-order.docx"
    document.save(source)

    first = compare_docx_documents(source, source)
    second = compare_docx_documents(source, source)
    script = (
        "import json, sys; from docxtor import compare_docx_documents; "
        "r=compare_docx_documents(sys.argv[1], sys.argv[1]); "
        "print(json.dumps([b.pair_id for b in r.right.blocks]))"
    )
    process = subprocess.run(
        [sys.executable, "-c", script, str(source)],
        check=True,
        capture_output=True,
        text=True,
        cwd=Path(__file__).resolve().parents[1],
        env={
            **os.environ,
            "PYTHONPATH": str(Path(__file__).resolve().parents[1] / "src"),
        },
    )
    subprocess_ids = json.loads(process.stdout)

    expected_locators = ["body:p:0", "header:0:p:0", "footer:0:p:0"]
    assert [block.coordinate.container_id for block in first.left.blocks] == expected_locators
    first_ids = [block.pair_id for block in first.left.blocks]
    assert first_ids == [block.pair_id for block in second.left.blocks]
    assert first_ids == subprocess_ids


def test_comment_anchor_offsets_follow_accepted_text_across_tracked_replacement() -> None:
    document = PyDocxDocument()
    paragraph = document.add_paragraph()
    paragraph.add_run("The notice period is 14 days.")
    document.add_comment(
        runs=paragraph.runs,
        text="Existing source comment.",
        author="Dike",
    )
    stream = BytesIO()
    document.save(stream)
    source = stream.getvalue()

    def mark_replacement(root) -> None:
        paragraph = root.xpath("//w:p", namespaces=_NS)[0]
        original_run = paragraph.find(f"{{{_W}}}r")
        assert original_run is not None
        run_index = list(paragraph).index(original_run)
        paragraph.remove(original_run)

        prefix = etree.Element(f"{{{_W}}}r")
        prefix_text = etree.SubElement(prefix, f"{{{_W}}}t")
        prefix_text.text = "The notice period is "
        prefix_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

        deletion = etree.Element(
            f"{{{_W}}}del",
            {
                f"{{{_W}}}id": "10",
                f"{{{_W}}}author": "Process reviewer",
            },
        )
        deleted_run = etree.SubElement(deletion, f"{{{_W}}}r")
        deleted_text = etree.SubElement(deleted_run, f"{{{_W}}}delText")
        deleted_text.text = "14"

        insertion = etree.Element(
            f"{{{_W}}}ins",
            {
                f"{{{_W}}}id": "11",
                f"{{{_W}}}author": "Process reviewer",
            },
        )
        inserted_run = etree.SubElement(insertion, f"{{{_W}}}r")
        inserted_text = etree.SubElement(inserted_run, f"{{{_W}}}t")
        inserted_text.text = "30"

        suffix = etree.Element(f"{{{_W}}}r")
        suffix_text = etree.SubElement(suffix, f"{{{_W}}}t")
        suffix_text.text = " days."
        suffix_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for offset, child in enumerate((prefix, deletion, insertion, suffix)):
            paragraph.insert(run_index + offset, child)

    reviewed = _rewrite_document(source, mark_replacement)

    result = compare_docx_documents(source, reviewed)

    assert not any(
        diagnostic.code == "comment_text_anchor_unresolved" for diagnostic in result.diagnostics
    )
    for view in (result.left, result.right):
        comment = view.comments[0]
        block = view.blocks[0]
        assert comment.author == "Dike"
        assert comment.anchor is not None
        assert comment.anchor.text == block.text
        assert comment.anchor.start_offset == 0
        assert comment.anchor.end_offset == len(block.text)


def test_comment_on_matched_block_is_stable_when_preceding_paragraph_is_inserted() -> None:
    document = PyDocxDocument()
    document.add_paragraph("Leading paragraph.")
    paragraph = document.add_paragraph("Commented paragraph.")
    document.add_comment(
        runs=paragraph.runs,
        text="Existing source comment.",
        author="Reviewer",
    )
    stream = BytesIO()
    document.save(stream)
    source = stream.getvalue()

    def insert_leading_paragraph(root) -> None:
        body = root.find(f"{{{_W}}}body")
        assert body is not None
        paragraph = etree.Element(f"{{{_W}}}p")
        run = etree.SubElement(paragraph, f"{{{_W}}}r")
        text = etree.SubElement(run, f"{{{_W}}}t")
        text.text = "A new leading paragraph."
        body.insert(0, paragraph)

    shifted = _rewrite_document(source, insert_leading_paragraph)

    result = compare_docx_documents(source, shifted)

    assert result.left.comments[0].anchor is not None
    assert result.right.comments[0].anchor is not None
    assert result.left.comments[0].anchor.locator != result.right.comments[0].anchor.locator
    assert result.left.comments[0].anchor.pair_id == result.right.comments[0].anchor.pair_id
    assert result.comment_changes == ()


def test_duplicate_revision_ids_are_not_silently_dropped() -> None:
    source = _docx(
        (
            "First tracked paragraph.",
            "Second tracked paragraph.",
            "Deleted tracked paragraph.",
        )
    )

    def reuse_revision_id(root) -> None:
        paragraphs = root.xpath("//w:body/w:p", namespaces=_NS)
        for paragraph in paragraphs[:2]:
            run = paragraph.find(f"{{{_W}}}r")
            assert run is not None
            index = list(paragraph).index(run)
            paragraph.remove(run)
            insertion = etree.Element(
                f"{{{_W}}}ins",
                {f"{{{_W}}}id": "9", f"{{{_W}}}author": "Reviewer"},
            )
            insertion.append(run)
            paragraph.insert(index, insertion)
        paragraph = paragraphs[2]
        run = paragraph.find(f"{{{_W}}}r")
        assert run is not None
        deleted_text = run.find(f"{{{_W}}}t")
        assert deleted_text is not None
        deleted_text.tag = f"{{{_W}}}delText"
        index = list(paragraph).index(run)
        paragraph.remove(run)
        deletion = etree.Element(
            f"{{{_W}}}del",
            {f"{{{_W}}}id": "9", f"{{{_W}}}author": "Reviewer"},
        )
        deletion.append(run)
        paragraph.insert(index, deletion)

    reviewed = _rewrite_document(source, reuse_revision_id)

    result = compare_docx_documents(source, reviewed)

    introduced = [event for event in result.revision_events if event.status.value == "introduced"]
    assert len(introduced) == 3
    assert {event.right.kind for event in introduced if event.right} == {"ins", "del"}
    assert len({event.right.pair_id for event in introduced if event.right}) == 3
    assert result.coverage.value == "incomplete"
    assert any(
        diagnostic.code == "duplicate_revision_identity" for diagnostic in result.diagnostics
    )


def test_inserted_paragraph_before_revision_does_not_change_revision_lifecycle() -> None:
    source = _docx(("Tracked paragraph.", "Following paragraph."))

    def mark_first_paragraph(root) -> None:
        paragraph = root.xpath("//w:body/w:p", namespaces=_NS)[0]
        run = paragraph.find(f"{{{_W}}}r")
        assert run is not None
        index = list(paragraph).index(run)
        paragraph.remove(run)
        insertion = etree.Element(
            f"{{{_W}}}ins",
            {f"{{{_W}}}id": "15", f"{{{_W}}}author": "Reviewer"},
        )
        insertion.append(run)
        paragraph.insert(index, insertion)

    tracked = _rewrite_document(source, mark_first_paragraph)

    def insert_leading_paragraph(root) -> None:
        body = root.find(f"{{{_W}}}body")
        assert body is not None
        paragraph = etree.Element(f"{{{_W}}}p")
        run = etree.SubElement(paragraph, f"{{{_W}}}r")
        text = etree.SubElement(run, f"{{{_W}}}t")
        text.text = "New leading paragraph."
        body.insert(0, paragraph)

    shifted = _rewrite_document(tracked, insert_leading_paragraph)

    result = compare_docx_documents(tracked, shifted)

    assert any(change.kind.value == "inserted" for change in result.text_changes)
    assert result.revision_events == ()


def test_revision_target_moving_within_same_paragraph_is_a_changed_event() -> None:
    def tracked(position: int) -> bytes:
        document = PyDocxDocument()
        paragraph = document.add_paragraph()
        for index, text in enumerate(("same", " middle ", "same")):
            run = paragraph.add_run(text)
            if index == position:
                insertion = etree.Element(
                    f"{{{_W}}}ins",
                    {f"{{{_W}}}id": "1", f"{{{_W}}}author": "Reviewer"},
                )
                run._r.getparent().replace(run._r, insertion)
                insertion.append(run._r)
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    result = compare_docx_documents(tracked(0), tracked(2))

    assert result.text_changes == ()
    assert len(result.revision_events) == 1
    event = result.revision_events[0]
    assert event.status.value == "changed"
    assert event.left is not None and event.left.start_offset == 0
    assert event.right is not None and event.right.start_offset == 12


def test_unsupported_revision_shapes_are_reported_as_incomplete_coverage() -> None:
    source = _rewrite_document(
        _docx(("Visible text.",)),
        lambda root: root.xpath("//w:p", namespaces=_NS)[0].append(
            etree.Element(f"{{{_W}}}moveFromRangeStart", {f"{{{_W}}}id": "5"})
        ),
    )

    result = compare_docx_documents(source, source)

    assert result.coverage.value == "incomplete"
    assert any(
        diagnostic.code == "unsupported_structural_revision" for diagnostic in result.diagnostics
    )


def test_unsupported_objects_are_reported_as_coverage_diagnostics() -> None:
    source = _rewrite_document(
        _docx(("Visible text.",)),
        lambda root: root.xpath("//w:p", namespaces=_NS)[0].append(
            etree.Element(f"{{{_W}}}object")
        ),
    )

    result = compare_docx_documents(source, source)

    assert result.coverage.value == "incomplete"
    assert any(diagnostic.code == "unsupported_object" for diagnostic in result.diagnostics)


def test_unsupported_visible_inline_character_is_reported_as_incomplete() -> None:
    left = _docx(("contractterm",))

    def add_no_break_hyphen(root) -> None:
        run = root.find(f".//{{{_W}}}r")
        assert run is not None
        etree.SubElement(run, f"{{{_W}}}noBreakHyphen")
        etree.SubElement(run, f"{{{_W}}}t").text = "term"

    right = _rewrite_document(_docx(("contract",)), add_no_break_hyphen)

    result = compare_docx_documents(left, right)

    assert result.text_changes == ()
    assert result.coverage.value == "incomplete"
    assert any(diagnostic.code == "unsupported_inline_surface" for diagnostic in result.diagnostics)


def test_style_inherited_numbering_has_an_unsupported_coverage_diagnostic() -> None:
    document = PyDocxDocument()
    document.add_paragraph("Numbered clause.", style="List Number")
    stream = BytesIO()
    document.save(stream)
    source = stream.getvalue()

    result = compare_docx_documents(source, source)

    assert result.coverage.value == "incomplete"
    assert any(diagnostic.code == "numbering_not_projected" for diagnostic in result.diagnostics)


def test_footnote_paragraph_projection_does_not_require_document_part_parent() -> None:
    left = _docx_with_footnote("Original footnote.")
    right = _docx_with_footnote("Corrected footnote.")

    result = compare_docx_documents(left, right)

    assert any(
        block.coordinate.story_kind == "footnote" and block.text == "Original footnote."
        for block in result.left.blocks
    )
    assert any(
        block.coordinate.story_kind == "footnote" and block.text == "Corrected footnote."
        for block in result.right.blocks
    )
    assert any(
        change.left is not None
        and change.right is not None
        and change.left.text == "Original"
        and change.right.text == "Corrected"
        for change in result.text_changes
    )
    assert result.coverage.value == "incomplete"
    assert any(diagnostic.code == "unsupported_inline_surface" for diagnostic in result.diagnostics)
