from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docxtor import DocxDocument, flatten_link_fields


@pytest.mark.parametrize("as_string", [False, True])
def test_unchanged_package_is_not_written(tmp_path: Path, as_string: bool) -> None:
    path = tmp_path / "unchanged.docx"
    document = Document()
    document.add_paragraph("Plain text without links.")
    document.save(path)
    before = path.read_bytes()
    stat = path.stat()

    assert flatten_link_fields(str(path) if as_string else path) is False

    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == stat.st_mtime_ns
    assert path.stat().st_ino == stat.st_ino


@pytest.mark.parametrize("story", ["document", "header", "footer"])
def test_rewritten_package_keeps_display_and_is_idempotent(tmp_path: Path, story: str) -> None:
    path = tmp_path / "links.docx"
    document = Document()
    if story == "document":
        paragraph = document.add_paragraph()
        part_name = "word/document.xml"
    else:
        container = getattr(document.sections[0], story)
        paragraph = container.paragraphs[0]
        part_name = f"word/{story}1.xml"
    paragraph.add_run("Contact: ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), ' HYPERLINK "mailto:field@example.com" ')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Field label"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    paragraph.add_run(" and ")
    hyperlink = OxmlElement("w:hyperlink")
    rid = paragraph.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Link label"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    document.save(path)
    before = path.read_bytes()

    assert flatten_link_fields(path) is True

    rewritten = path.read_bytes()
    assert rewritten != before
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read(part_name))
        assert "".join(node.text or "" for node in root.iter(qn("w:t"))) == (
            "Contact: Field label and Link label"
        )
        assert not list(root.iter(qn("w:fldSimple")))
        assert not list(root.iter(qn("w:hyperlink")))
        rels_name = part_name.replace("word/", "word/_rels/") + ".rels"
        relationships = ET.fromstring(archive.read(rels_name))
        assert all(rel.get("Id") != rid for rel in relationships)
    reopened = DocxDocument.open(path)
    assert any(
        segment.text == "Contact: Field label and Link label" for segment in reopened.segments
    )

    assert flatten_link_fields(path) is False
    assert path.read_bytes() == rewritten


@pytest.mark.parametrize("part_name", ["word/document.xml", "word/header1.xml", "word/footer1.xml"])
def test_malformed_story_preserves_source(tmp_path: Path, part_name: str) -> None:
    path = tmp_path / "malformed.docx"
    document = Document()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Header"
    document.sections[0].footer.paragraphs[0].text = "Footer"
    data = BytesIO()
    document.save(data)
    with ZipFile(data) as source, ZipFile(path, "w") as output:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == part_name:
                payload = b"<broken"
            output.writestr(info, payload)
    before = path.read_bytes()

    with pytest.raises(ET.ParseError):
        flatten_link_fields(path)

    assert path.read_bytes() == before
    assert set(tmp_path.iterdir()) == {path}
