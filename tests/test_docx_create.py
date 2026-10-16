from io import BytesIO
from pathlib import Path

from docx import Document

from docxtor import create_docx_from_paragraphs
from docxtor.docx_create import write_docx_from_paragraphs


def test_create_docx_from_paragraphs(tmp_path: Path) -> None:
    data = create_docx_from_paragraphs(("one", "two"))
    assert [item.text for item in Document(BytesIO(data)).paragraphs] == ["one", "two"]
    path = tmp_path / "created.docx"
    assert write_docx_from_paragraphs(path, ("one", "two")) == path
    assert [item.text for item in Document(path).paragraphs] == ["one", "two"]
