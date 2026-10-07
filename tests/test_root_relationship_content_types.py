from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from docxtor import (
    DocumentMark,
    DocxDocument,
    PhysicalReviewEdit,
    docx_facts,
    inspect_revision_dispositions,
    inventory_review_markup,
    remove_all_comments_bytes,
    render_physical_clean,
    stamp_document_mark,
)

RELATIONSHIP_TYPE = "application/vnd.openxmlformats-package.relationships+xml"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"


def document_bytes(*, comments=False):
    document = Document()
    paragraph = document.add_paragraph("Term is 30 days.")
    paragraph.runs[0].bold = True
    document.add_paragraph("Unchanged.")
    document.sections[0].footer.paragraphs[0].text = "Original footer"
    if comments:
        document.add_comment(paragraph.runs[0], text="Use 14 days.", author="Synthetic")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def assert_complete(data):
    part = next(part for part in docx_facts(data).parts if part.name == "_rels/.rels")
    assert part.content_type == RELATIONSHIP_TYPE
    assert part.is_xml
    assert DocxDocument.open_bytes(data).inventory().coverage.value == "complete"
    review = inventory_review_markup(data)
    assert review.coverage.value == "complete"
    assert not review.diagnostics
    revision = inspect_revision_dispositions(data)
    assert revision.input_sha256 == sha256(data).hexdigest()
    assert revision.coverage.value == "complete"
    assert revision.count == 0
    assert not revision.diagnostics


@pytest.mark.parametrize("comments", [False, True])
def test_default_root_relationship_type_has_complete_public_coverage(comments):
    data = document_bytes(comments=comments)
    with ZipFile(BytesIO(data)) as archive:
        types = etree.fromstring(archive.read("[Content_Types].xml"))
        assert not any(child.get("PartName") == "/_rels/.rels" for child in types)
    assert_complete(data)


def test_clean_render_and_publication_preserve_complete_default_relationship_coverage(tmp_path):
    source = document_bytes(comments=True)
    cleaned = remove_all_comments_bytes(source)
    assert cleaned.receipt.before_sha256 == sha256(source).hexdigest()
    assert_complete(cleaned.data)
    output = tmp_path / "clean.docx"
    render_physical_clean(
        cleaned.data,
        output,
        (
            PhysicalReviewEdit(
                action_id="term",
                locator="body:p:0",
                operation="replace",
                start_offset=0,
                end_offset=len("Term is 30 days."),
                expected_text="Term is 30 days.",
                replacement_text="Term is 14 days.",
            ),
        ),
    )
    marked = stamp_document_mark(
        output.read_bytes(), DocumentMark("clean", "Synthetic", ("clean",), "Publication mark")
    )
    assert_complete(marked)
    assert not inventory_review_markup(marked).comments
    document = Document(BytesIO(marked))
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Term is 14 days.",
        "Unchanged.",
    ]
    assert all(run.bold for run in document.paragraphs[0].runs if run.text)
    assert "Original footer" in "\n".join(p.text for p in document.sections[0].footer.paragraphs)


@pytest.mark.parametrize("mutation", ["missing-default", "override", "unknown-part"])
def test_explicit_overrides_are_respected_and_unknown_parts_stay_incomplete(mutation):
    stream = BytesIO()
    with ZipFile(BytesIO(document_bytes())) as original, ZipFile(stream, "w") as modified:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "[Content_Types].xml":
                root = etree.fromstring(data)
                if mutation == "missing-default":
                    for child in list(root):
                        if child.get("Extension") == "rels":
                            root.remove(child)
                elif mutation == "override":
                    etree.SubElement(
                        root,
                        CT + "Override",
                        PartName="/_rels/.rels",
                        ContentType="application/octet-stream",
                    )
                data = etree.tostring(root)
            modified.writestr(item, data)
        if mutation == "unknown-part":
            modified.writestr("word/unregistered.mystery", b"opaque")
    data = stream.getvalue()
    if mutation == "override":
        part = next(part for part in docx_facts(data).parts if part.name == "_rels/.rels")
        assert part.content_type == "application/octet-stream"
        return
    inventory = inspect_revision_dispositions(data)
    assert inventory.coverage.value == "incomplete"
    assert any(row.code == "part_type_unknown" for row in inventory.diagnostics)
    if mutation != "unknown-part":
        part = next(part for part in docx_facts(data).parts if part.name == "_rels/.rels")
        assert part.content_type == "application/octet-stream"
