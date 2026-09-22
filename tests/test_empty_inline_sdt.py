"""Zero-width SDTs survive text replacement without admitting other opaque content."""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docxtor import (
    PhysicalReviewEdit,
    PhysicalReviewPlan,
    accept_all_revisions_bytes,
    reject_all_revisions_bytes,
    render_physical_clean,
    render_physical_review,
)
from docxtor.docx_review_render import PhysicalReviewRenderError


def source_bytes(*, payload=None, positions=(7,)):
    doc = Document()
    p = doc.add_paragraph()
    text = "Termin 10 dni."
    last = 0
    for index, position in enumerate(positions):
        p.add_run(text[last:position]).bold = True
        sdt = OxmlElement("w:sdt")
        pr = OxmlElement("w:sdtPr")
        tag = OxmlElement("w:tag")
        tag.set(qn("w:val"), f"marker-{index}")
        pr.append(tag)
        sdt.append(pr)
        content = OxmlElement("w:sdtContent")
        if payload is not None:
            content.append(OxmlElement(payload))
        sdt.append(content)
        p._p.append(sdt)
        last = position
    p.add_run(text[last:]).bold = True
    doc.add_paragraph("Untouched.")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def assert_document(data, text, count):
    doc = Document(BytesIO(data))
    assert [p.text for p in doc.paragraphs] == [text, "Untouched."]
    assert [t.get(qn("w:val")) for t in doc.paragraphs[0]._p.iter(qn("w:tag"))] == [
        f"marker-{i}" for i in range(count)
    ]
    assert all(r.bold for r in doc.paragraphs[0].runs if r.text)


@pytest.mark.parametrize("tracked", [True, False])
@pytest.mark.parametrize("positions", [(7,), (0,), (14,), (3, 7, 11)])
@pytest.mark.parametrize("replacement", ["Termin 30 dni.", "Krótko.", ""])
def test_empty_sdt_survives_replacement(tmp_path, tracked, positions, replacement):
    data = source_bytes(positions=positions)
    output = tmp_path / "output.docx"
    edit = PhysicalReviewEdit(
        action_id="change",
        locator="body:p:0",
        operation="replace",
        start_offset=0,
        end_offset=14,
        expected_text="Termin 10 dni.",
        replacement_text=replacement,
    )
    if tracked:
        render_physical_review(data, output, PhysicalReviewPlan(edits=(edit,)))
        accepted = accept_all_revisions_bytes(output.read_bytes()).output_bytes
        rejected = reject_all_revisions_bytes(output.read_bytes()).output_bytes
        assert_document(accepted, replacement, len(positions))
        assert_document(rejected, "Termin 10 dni.", len(positions))
        # Reject restores each marker's position among the original text runs.
        p = Document(BytesIO(rejected)).paragraphs[0]._p
        pos = 0
        found = []
        for child in p:
            if child.tag == qn("w:sdt"):
                found.append(pos)
            else:
                pos += sum(len(t.text or "") for t in child.iter(qn("w:t")))
        assert found == list(positions)
    else:
        render_physical_clean(data, output, (edit,))
        assert_document(output.read_bytes(), replacement, len(positions))


@pytest.mark.parametrize("tracked", [True, False])
@pytest.mark.parametrize("payload", ["w:r", "w:drawing", "w:fldSimple"])
def test_nonempty_sdt_is_still_protected(tmp_path, tracked, payload):
    data = source_bytes(payload=payload)
    output = tmp_path / "output.docx"
    edit = PhysicalReviewEdit(
        action_id="change",
        locator="body:p:0",
        operation="replace",
        start_offset=0,
        end_offset=14,
        expected_text="Termin 10 dni.",
        replacement_text="Changed.",
    )
    with pytest.raises(PhysicalReviewRenderError, match="opaque content"):
        if tracked:
            render_physical_review(data, output, PhysicalReviewPlan(edits=(edit,)))
        else:
            render_physical_clean(data, output, (edit,))
    assert not output.exists()
