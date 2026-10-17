from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docxtor import (
    AcceptRevisionsError,
    CommentAuthor,
    CommentRange,
    DocxDocument,
    RevisionInventoryCoverage,
    accept_all_revisions_bytes,
    inventory_review_markup,
    inventory_revisions_bytes,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _document(value: str | None = None) -> bytes:
    document = Document()
    paragraph = document.add_paragraph("First ")
    paragraph.runs[0].italic = True
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "8")
    inserted.set(qn("w:author"), "Reviewer")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "accepted sentence."
    run.append(text)
    inserted.append(run)
    paragraph._p.append(inserted)
    setting = OxmlElement("w:trackRevisions")
    if value is not None:
        setting.set(qn("w:val"), value)
    document.settings.element.append(setting)
    for name in ("doNotTrackMoves", "doNotTrackFormatting"):
        document.settings.element.append(OxmlElement(f"w:{name}"))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _parts(data: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(data)) as package:
        return {name: package.read(name) for name in package.namelist()}


def _replace_parts(data: bytes, replacements: dict[str, bytes]) -> bytes:
    parts = _parts(data) | replacements
    output = BytesIO()
    with ZipFile(output, "w") as package:
        for name, payload in parts.items():
            package.writestr(name, payload)
    return output.getvalue()


@pytest.mark.parametrize(
    ("value", "enabled"),
    [
        (None, True),
        ("true", True),
        ("on", True),
        ("1", True),
        ("false", False),
        ("off", False),
        ("0", False),
    ],
)
def test_inventory_reads_tracking_and_default_acceptance_preserves_setting(value, enabled) -> None:
    source = _document(value)
    inventory = inventory_revisions_bytes(source)
    result = accept_all_revisions_bytes(source, drop_comments=False)
    assert inventory.coverage is RevisionInventoryCoverage.COMPLETE
    assert inventory.tracking_enabled is enabled
    assert result.before.tracking_enabled is enabled
    assert result.after.tracking_enabled is enabled
    assert result.after.count == 0
    assert _parts(source)["word/settings.xml"] == _parts(result.output_bytes)["word/settings.xml"]
    assert Document(BytesIO(result.output_bytes)).paragraphs[0].text == "First accepted sentence."


def test_acceptance_can_disable_tracking_without_changing_other_settings_or_parts() -> None:
    source = _document("true")
    result = accept_all_revisions_bytes(source, drop_comments=False, disable_tracking=True)
    assert result.before.count == 1 and result.before.tracking_enabled is True
    assert result.after.count == 0 and result.after.tracking_enabled is False
    assert result.after.coverage is RevisionInventoryCoverage.COMPLETE
    assert result.after.diagnostics == ()
    before, after = _parts(source), _parts(result.output_bytes)
    assert before.keys() == after.keys()
    for name in before.keys() - {"word/document.xml", "word/settings.xml"}:
        assert before[name] == after[name], name
    expected = etree.fromstring(before["word/settings.xml"])
    expected.remove(expected.find(qn("w:trackRevisions")))
    actual = etree.fromstring(after["word/settings.xml"])
    assert etree.tostring(expected, method="c14n") == etree.tostring(actual, method="c14n")
    document = Document(BytesIO(result.output_bytes))
    assert document.paragraphs[0].text == "First accepted sentence."
    assert document.paragraphs[0].runs[0].italic is True
    again = accept_all_revisions_bytes(
        result.output_bytes, drop_comments=False, disable_tracking=True
    )
    assert _parts(again.output_bytes) == after


def test_tracking_cleanup_covers_orphan_settings_with_an_arbitrary_part_name() -> None:
    source = _document("off")
    source = _replace_parts(
        source,
        {
            "customXml/orphan-settings.dat": f'<w:settings xmlns:w="{W}">'
            '<w:trackRevisions w:val="on"/></w:settings>'.encode()
        },
    )
    assert inventory_revisions_bytes(source).tracking_enabled is True
    result = accept_all_revisions_bytes(source, drop_comments=False, disable_tracking=True)
    assert result.before.tracking_enabled is True
    assert result.after.tracking_enabled is False
    root = etree.fromstring(_parts(result.output_bytes)["customXml/orphan-settings.dat"])
    assert root.find(qn("w:trackRevisions")) is None


@pytest.mark.parametrize(
    "settings",
    [
        '<w:settings><w:trackRevisions w:val="maybe"/></w:settings>',
        "<w:settings><w:trackRevisions/><w:trackRevisions/></w:settings>",
        "<w:settings><w:trackRevisions><w:r/></w:trackRevisions></w:settings>",
        "<w:document><w:trackRevisions/></w:document>",
    ],
)
def test_acceptance_refuses_invalid_tracking_settings(settings: str) -> None:
    source = _replace_parts(
        _document(),
        {"word/settings.xml": settings.replace(">", f' xmlns:w="{W}">', 1).encode()},
    )
    inventory = inventory_revisions_bytes(source)
    assert inventory.coverage is RevisionInventoryCoverage.INCOMPLETE
    with pytest.raises(AcceptRevisionsError, match="inventory is incomplete"):
        accept_all_revisions_bytes(source, drop_comments=False, disable_tracking=True)


def test_acceptance_refuses_malformed_formatting_snapshots_before_erasing_them() -> None:
    source = _document()
    root = etree.fromstring(_parts(source)["word/document.xml"])
    run = next(root.iter(qn("w:r")))
    properties, change = OxmlElement("w:rPr"), OxmlElement("w:rPrChange")
    change.append(OxmlElement("w:pPr"))
    properties.append(change)
    run.insert(0, properties)
    source = _replace_parts(source, {"word/document.xml": etree.tostring(root)})
    assert inventory_revisions_bytes(source).coverage is RevisionInventoryCoverage.INCOMPLETE
    with pytest.raises(AcceptRevisionsError, match="malformed rPrChange"):
        accept_all_revisions_bytes(source, drop_comments=False, disable_tracking=True)


def test_disabling_tracking_does_not_resolve_comments() -> None:
    source = (
        DocxDocument.open_bytes(_document())
        .add_comment(
            CommentRange("body:p:0", 0, 5, expected_text="First"),
            "Unresolved instruction",
            CommentAuthor("Reviewer"),
        )
        .data
    )
    result = accept_all_revisions_bytes(source, drop_comments=False, disable_tracking=True)
    comments = inventory_review_markup(result.output_bytes).comments
    assert len(comments) == 1
    assert comments[0].text == "Unresolved instruction"
    assert result.after.tracking_enabled is False
