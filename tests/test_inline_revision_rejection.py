from __future__ import annotations

import time
import zipfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docxtor import (
    DocumentError,
    DocxDocument,
    InlineRevisionGroup,
    InlineSegment,
    PhysicalReviewEdit,
    PhysicalReviewPlan,
    group_inline_revisions,
    paragraph_to_inline_segments,
    rebuild_paragraph_from_inline,
    render_physical_review,
    restore_deleted_inline,
)


@pytest.fixture(autouse=True)
def deterministic_archive_clock(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep exact package-byte guards independent of ZIP's wall-clock timestamps.

    python-docx serializes through ZipFile.writestr, which stamps each entry from
    zipfile.time. Freeze only that module's clock facade, not document payloads
    or byte comparisons; a public inline mutation must still change the bytes.
    """
    archive_time = time.localtime(1760000000.0)
    monkeypatch.setattr(
        zipfile,
        "time",
        SimpleNamespace(time=lambda: time.time(), localtime=lambda _seconds: archive_time),
    )


def _revision(
    kind: str,
    text: str,
    *,
    identity: str | None = "10",
    author: str = "A",
    date: str = "2026-01-01T00:00:00Z",
    italic: bool = False,
):
    wrapper = OxmlElement(f"w:{kind}")
    if identity is not None:
        wrapper.set(qn("w:id"), identity)
    wrapper.set(qn("w:author"), author)
    wrapper.set(qn("w:date"), date)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    props.append(OxmlElement("w:i" if italic else "w:b"))
    run.append(props)
    node = OxmlElement("w:delText" if kind == "del" else "w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    wrapper.append(run)
    return wrapper


def _xml(node):
    return etree.tostring(node, method="c14n", exclusive=True)


def _bytes(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_split_deletion_restores_each_run_and_preserves_source(tmp_path: Path) -> None:
    source = Document()
    para = source.add_paragraph()
    para._p.append(_revision("del", " Left "))
    para._p.append(_revision("del", "Right ", italic=True))
    para._p.append(_revision("ins", "Applied", identity="12"))
    data = _bytes(source)
    source_path = tmp_path / "source.docx"
    source_path.write_bytes(data)
    document = DocxDocument.open(source_path)
    groups = group_inline_revisions(document.get_inline_segments("body:p:0"))
    assert [group.text for group in groups] == [" Left Right ", "Applied"]
    assert [group.kind for group in groups] == ["del", "ins"]
    xml_before = [_xml(seg.element) for seg in groups[0].segments]
    restored = restore_deleted_inline(groups[0])
    assert [item.text for item in restored] == [" Left ", "Right "]
    assert [_xml(seg.element) for seg in groups[0].segments] == xml_before
    rebuild_paragraph_from_inline(
        document.resolve_paragraph("body:p:0"), [*restored, *groups[1].segments]
    )
    output = tmp_path / "restored.docx"
    document.publish(output)
    result = Document(output)
    assert result.paragraphs[0].runs[0].bold
    assert result.paragraphs[0].runs[1].italic
    assert [n.text for n in result.element.iter(qn("w:t"))] == [" Left ", "Right ", "Applied"]
    assert len(list(result.element.iter(qn("w:ins")))) == 1
    assert source_path.read_bytes() == data


@pytest.mark.parametrize("boundary", ["id", "author", "date", "insertion", "plain", "missing"])
def test_grouping_respects_physical_identity_boundaries(boundary: str) -> None:
    source = Document()
    para = source.add_paragraph()
    para._p.append(_revision("del", "Left", identity=None if boundary == "missing" else "10"))
    if boundary == "plain":
        para.add_run("Middle")
    elif boundary == "insertion":
        para._p.append(_revision("ins", "Middle"))
    para._p.append(
        _revision(
            "del",
            "Right",
            identity="11" if boundary == "id" else None if boundary == "missing" else "10",
            author="B" if boundary == "author" else "A",
            date="2026-01-02T00:00:00Z" if boundary == "date" else "2026-01-01T00:00:00Z",
        )
    )
    groups = group_inline_revisions(paragraph_to_inline_segments(para))
    assert [group.text for group in groups] == (
        ["Left", "Middle", "Right"] if boundary in {"plain", "insertion"} else ["Left", "Right"]
    )


@pytest.mark.parametrize("opaque_kind", ["sdt", "fldSimple"])
@pytest.mark.parametrize("nested_kind", [None, "del", "ins", "moveFrom", "moveTo"])
def test_opaque_restoration_keeps_nested_revisions_pending(opaque_kind, nested_kind) -> None:
    source = Document()
    para = source.add_paragraph()
    outer = _revision("del", "")
    for child in list(outer):
        outer.remove(child)
    opaque = OxmlElement(f"w:{opaque_kind}")
    if opaque_kind == "sdt":
        props = OxmlElement("w:sdtPr")
        tag = OxmlElement("w:tag")
        tag.set(qn("w:val"), "preserved")
        props.append(tag)
        opaque.append(props)
        content = OxmlElement("w:sdtContent")
        opaque.append(content)
    else:
        opaque.set(qn("w:instr"), " MERGEFIELD preserved ")
        content = opaque
    content.append(_revision("del", " Left ")[0])
    nested = _revision(nested_kind, "Pending", identity="11") if nested_kind else None
    if nested is not None:
        content.append(nested)
    content.append(_revision("del", "Right ", italic=True)[0])
    outer.append(opaque)
    para._p.append(outer)
    before = _xml(outer)
    group = group_inline_revisions(paragraph_to_inline_segments(para))[0]
    restored = restore_deleted_inline(group)
    assert _xml(outer) == before
    retained = restored[0].element
    assert retained.tag == qn(f"w:{opaque_kind}")
    if opaque_kind == "sdt":
        assert _xml(retained.find(qn("w:sdtPr"))) == _xml(opaque.find(qn("w:sdtPr")))
    else:
        assert retained.get(qn("w:instr")) == opaque.get(qn("w:instr"))
    assert "".join(item.text for item in restored) == group.text
    if nested is not None:
        assert _xml(next(retained.iter(qn(f"w:{nested_kind}")))) == _xml(nested)
    nodes = (
        [
            n
            for n in retained.iter(qn("w:t"))
            if not any(p.tag == qn(f"w:{nested_kind}") for p in n.iterancestors())
        ]
        if nested_kind
        else list(retained.iter(qn("w:t")))
    )
    assert [n.text for n in nodes] == [" Left ", "Right "]
    assert all(n.get(qn("xml:space")) == "preserve" for n in nodes)
    expected_props = [_xml(content[0].find(qn("w:rPr"))), _xml(content[-1].find(qn("w:rPr")))]
    direct = retained.find(qn("w:sdtContent")) if opaque_kind == "sdt" else retained
    assert [_xml(r.find(qn("w:rPr"))) for r in direct if r.tag == qn("w:r")] == expected_props


@pytest.mark.parametrize("invalid", ["text", "insertion", "missing", "mismatch", "empty"])
def test_restoration_refuses_invalid_group_without_mutation(invalid: str) -> None:
    node = _revision("ins" if invalid == "insertion" else "del", "Old")
    segment = InlineSegment(
        "text" if invalid == "text" else "opaque",
        "Old",
        element=None if invalid == "missing" else node,
    )
    group = InlineRevisionGroup(
        segments=() if invalid == "empty" else (segment,),
        kind="ins" if invalid == "insertion" else "del",
        text="Wrong" if invalid == "mismatch" else "Old",
    )
    before = _xml(node)
    with pytest.raises(DocumentError):
        restore_deleted_inline(group)
    assert _xml(node) == before


def test_remove_inserted_paragraph_refreshes_addresses_and_preserves_source(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("Before")
    inserted = source.add_paragraph()
    inserted._p.append(_revision("ins", "Inserted"))
    source.add_paragraph("After")
    data = _bytes(source)
    source_path = tmp_path / "source.docx"
    source_path.write_bytes(data)
    document = DocxDocument.open(source_path)
    document.remove_inserted_paragraph("body:p:1", expected_text="Inserted")
    assert [p.value for p in document.paragraph_resolutions] == ["Before", "After"]
    assert [p.identity.container_id for p in document.paragraph_resolutions] == [
        "body:p:0",
        "body:p:1",
    ]
    assert document.texts == ["Before", "After"]
    assert all(span.role == "run" for span in document.spans)
    output = tmp_path / "removed.docx"
    document.publish(output)
    assert [p.text for p in Document(output).paragraphs] == ["Before", "After"]
    assert source_path.read_bytes() == data


@pytest.mark.parametrize("invalid", ["missing", "stale", "mixed", "opaque"])
def test_remove_inserted_paragraph_refuses_unowned_payload(invalid: str) -> None:
    source = Document()
    para = source.add_paragraph()
    para._p.append(_revision("ins", "Inserted"))
    if invalid == "mixed":
        para.add_run("Source")
    if invalid == "opaque":
        para._p.append(OxmlElement("w:sdt"))
    document = DocxDocument.open_bytes(_bytes(source))
    before = document.to_bytes()
    with pytest.raises(DocumentError):
        document.remove_inserted_paragraph(
            "body:p:99" if invalid == "missing" else "body:p:0",
            expected_text="Wrong" if invalid == "stale" else "Inserted",
        )
    assert document.to_bytes() == before


def test_same_identity_in_separate_paragraphs_stays_separate() -> None:
    source = Document()
    paragraphs = [source.add_paragraph(), source.add_paragraph()]
    for paragraph, text in zip(paragraphs, ("Left", "Right"), strict=True):
        paragraph._p.append(_revision("del", text))
    groups = [group_inline_revisions(paragraph_to_inline_segments(p)) for p in paragraphs]
    assert [[group.text for group in paragraph] for paragraph in groups] == [["Left"], ["Right"]]


def test_text_mismatch_inside_payload_fails_closed() -> None:
    payload = _revision("del", "Actual")
    segment = InlineSegment("opaque", "Different", element=payload)
    before = _xml(payload)
    with pytest.raises(DocumentError, match="restored deleted inline text"):
        restore_deleted_inline(InlineRevisionGroup((segment,), "del", "Different"))
    assert _xml(payload) == before


def test_remove_renderer_inserted_paragraph_preserves_original_package(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("Anchor")
    source_path = tmp_path / "source.docx"
    source_bytes = _bytes(source)
    source_path.write_bytes(source_bytes)
    reviewed = tmp_path / "reviewed.docx"
    render_physical_review(
        source_path,
        reviewed,
        PhysicalReviewPlan(
            edits=(
                PhysicalReviewEdit(
                    action_id="physical-insertion",
                    locator="body:p:0",
                    operation="insert",
                    start_offset=0,
                    end_offset=0,
                    replacement_text="Inserted",
                    new_paragraph=True,
                ),
            )
        ),
    )
    reviewed_bytes = reviewed.read_bytes()
    handle = DocxDocument.open(reviewed)
    handle.remove_inserted_paragraph("body:p:1", expected_text="Inserted")
    assert handle.texts == ["Anchor"]
    assert handle._source_bytes == reviewed_bytes
    handle.publish(tmp_path / "output.docx")
    assert reviewed.read_bytes() == reviewed_bytes
    assert source_path.read_bytes() == source_bytes


def test_public_removal_method_has_typed_keyword_guard_and_no_extra_export() -> None:
    import inspect

    import docxtor

    signature = inspect.signature(DocxDocument.remove_inserted_paragraph)
    assert str(signature) == "(self, container_id: 'str', *, expected_text: 'str') -> 'None'"
    assert "remove_inserted_paragraph" in dir(DocxDocument)
    assert "DocxParagraphMutationOperations" not in dir(docxtor)


def test_removal_refuses_table_story_without_mutating_package() -> None:
    source = Document()
    paragraph = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph._p.append(_revision("ins", "Inserted"))
    handle = DocxDocument.open_bytes(_bytes(source))
    locator = next(s.container_id for s in handle.spans if s.text == "Inserted")
    before = handle.to_bytes()
    with pytest.raises(DocumentError, match="supported body story"):
        handle.remove_inserted_paragraph(locator, expected_text="Inserted")
    assert handle.to_bytes() == before


def test_removal_refuses_inserted_paragraph_carrying_source_section_properties() -> None:
    source = Document()
    source.add_paragraph("Before")
    source.add_section()
    paragraph = source.paragraphs[-1]
    paragraph._p.append(_revision("ins", "Inserted"))
    source.add_paragraph("After")
    source_bytes = _bytes(source)
    handle = DocxDocument.open_bytes(source_bytes)
    before = handle.to_bytes()
    with pytest.raises(DocumentError, match="section"):
        handle.remove_inserted_paragraph("body:p:1", expected_text="Inserted")
    assert handle.to_bytes() == before
    assert handle._source_bytes == source_bytes
    assert len(Document(BytesIO(handle.to_bytes())).sections) == 2


def test_removal_checks_live_text_after_public_inline_rebuild() -> None:
    source = Document()
    source.add_paragraph()._p.append(_revision("ins", "Original insertion"))
    source.add_paragraph("After")
    handle = DocxDocument.open_bytes(_bytes(source))
    replacement = Document().add_paragraph()
    replacement._p.append(_revision("ins", "Changed insertion"))
    rebuild_paragraph_from_inline(
        handle.resolve_paragraph("body:p:0"), paragraph_to_inline_segments(replacement)
    )
    before = handle.to_bytes()
    with pytest.raises(DocumentError, match="expected entirely inserted"):
        handle.remove_inserted_paragraph("body:p:0", expected_text="Original insertion")
    assert handle.to_bytes() == before
    assert "".join(s.text for s in handle.get_inline_segments("body:p:0")) == "Changed insertion"
    handle.remove_inserted_paragraph("body:p:0", expected_text="Changed insertion")
    assert handle.texts == ["After"]


def test_removal_refuses_cached_paragraph_attached_to_foreign_body() -> None:
    source = Document()
    source.add_paragraph()._p.append(_revision("ins", "Inserted"))
    source.add_paragraph("After")
    handle = DocxDocument.open_bytes(_bytes(source))
    paragraph = handle.resolve_paragraph("body:p:0")
    paragraph._p.getparent().remove(paragraph._p)
    foreign_body = OxmlElement("w:body")
    foreign_body.append(paragraph._p)
    before = handle.to_bytes()
    foreign_before = _xml(foreign_body)
    with pytest.raises(DocumentError, match="supported body story"):
        handle.remove_inserted_paragraph("body:p:0", expected_text="Inserted")
    assert handle.to_bytes() == before
    assert _xml(foreign_body) == foreign_before


def test_removal_refuses_cached_locator_after_same_body_paragraph_reorder() -> None:
    source = Document()
    source.add_paragraph()._p.append(_revision("ins", "Inserted"))
    source.add_paragraph("After")
    handle = DocxDocument.open_bytes(_bytes(source))
    paragraph = handle.resolve_paragraph("body:p:0")
    after = handle.resolve_paragraph("body:p:1")
    after._p.addnext(paragraph._p)
    before = handle.to_bytes()
    cached = handle.texts
    with pytest.raises(DocumentError, match="same paragraph"):
        handle.remove_inserted_paragraph("body:p:0", expected_text="Inserted")
    assert handle.to_bytes() == before
    assert handle.texts == cached


def test_exact_bytes_detect_payload_changes_across_wall_clock_rollover(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    wall_clock = [1760000000.0]
    monkeypatch.setattr(time, "time", lambda: wall_clock[0])
    source = Document()
    source.add_paragraph("Original")
    handle = DocxDocument.open_bytes(_bytes(source))
    before = handle.to_bytes()
    wall_clock[0] += 2  # Cross the ZIP/DOS two-second timestamp boundary without any edit.
    assert handle.to_bytes() == before
    replacement = Document().add_paragraph("Changed")
    rebuild_paragraph_from_inline(
        handle.resolve_paragraph("body:p:0"), paragraph_to_inline_segments(replacement)
    )
    assert handle.to_bytes() != before
    assert "".join(s.text for s in handle.get_inline_segments("body:p:0")) == "Changed"
