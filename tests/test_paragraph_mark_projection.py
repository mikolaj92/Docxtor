from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docxtor import project_docx_for_review


def test_paragraph_marks_have_addressable_zero_width_revision_receipts(tmp_path):
    doc = Document()
    for i, kind in enumerate(("ins", "del")):
        paragraph = doc.add_paragraph(f"Body {i}.")
        properties = OxmlElement("w:rPr")
        mark = OxmlElement(f"w:{kind}")
        mark.set(qn("w:id"), str(i + 10))
        mark.set(qn("w:author"), "Synthetic")
        properties.append(mark)
        paragraph._p.get_or_add_pPr().append(properties)
    path = tmp_path / "marks.docx"
    doc.save(path)
    projection = project_docx_for_review(path)
    marks = getattr(projection, "paragraph_mark_revisions", ())
    assert len(marks) == 2
    assert [(m.container_id, m.role, m.text, m.start_offset, m.end_offset) for m in marks] == [
        ("body:p:0", "insertion", "", 7, 7),
        ("body:p:1", "deletion", "", 7, 7),
    ]
    assert projection.coverage.value == "incomplete"
    assert any(d.code == "unprojected_paragraph_merge" for d in projection.diagnostics)
    assert [m.revision_id for m in marks] == ["10", "11"]
    assert [p.text for p in projection.paragraphs] == ["Body 0.", "Body 1."]


def test_deleted_mark_on_empty_paragraph_is_not_lost_from_projection(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph()
    properties = OxmlElement("w:rPr")
    mark = OxmlElement("w:del")
    mark.set(qn("w:id"), "12")
    properties.append(mark)
    paragraph._p.get_or_add_pPr().append(properties)
    doc.add_paragraph("Retained.")
    path = tmp_path / "empty-mark.docx"
    doc.save(path)
    projection = project_docx_for_review(path)
    assert projection.coverage.value == "complete"
    assert len(projection.paragraph_mark_revisions) == 1
    assert projection.paragraph_mark_revisions[0].container_id == "body:p:0"
    assert projection.paragraph_mark_revisions[0].start_offset == 0


def test_paragraph_format_history_is_not_a_missing_boundary_revision(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Retained.")
    properties = OxmlElement("w:rPr")
    change = OxmlElement("w:rPrChange")
    change.set(qn("w:id"), "13")
    change.append(OxmlElement("w:rPr"))
    properties.append(change)
    paragraph._p.get_or_add_pPr().append(properties)
    path = tmp_path / "format-history.docx"
    doc.save(path)
    projection = project_docx_for_review(path)
    assert projection.coverage.value == "complete"
    assert projection.paragraph_mark_revisions == ()
