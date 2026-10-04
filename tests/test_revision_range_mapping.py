from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document as PyDocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docxtor import (
    CommentAuthor,
    CommentRange,
    DocxDocument,
    RevisionAuthor,
    RevisionMutationError,
    RevisionPosition,
    RevisionRange,
)
from docxtor.docx_review_models import OperationStatus
from docxtor.docx_revisions import accept_all_revisions_bytes

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
AUTHOR = CommentAuthor("Reviewer", "RV", "2024-01-01T00:00:00Z")
REVIEWER = RevisionAuthor("Reviewer", "2024-01-01T00:00:00Z")


def _plain_docx(*paragraphs: str) -> bytes:
    texts = paragraphs or ("Hello world",)
    document = PyDocxDocument()
    document.add_paragraph(texts[0])
    for text in texts[1:]:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _part_xml(data: bytes, part: str = "word/document.xml") -> etree._Element:
    with ZipFile(BytesIO(data)) as archive:
        return etree.fromstring(archive.read(part))


def _visible_text(data: bytes) -> str:
    return "".join(_part_xml(data).xpath("//w:t/text()", namespaces=NS))


def _del_text(data: bytes) -> list[str]:
    return [str(value) for value in _part_xml(data).xpath("//w:delText/text()", namespaces=NS)]


def _ins_text(data: bytes) -> list[str]:
    return [str(value) for value in _part_xml(data).xpath("//w:ins//w:t/text()", namespaces=NS)]


def _write_docx(path: Path, document: PyDocxDocument) -> None:
    buffer = BytesIO()
    document.save(buffer)
    path.write_bytes(buffer.getvalue())


def _append_run(parent: object, text: str) -> None:
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    parent.append(run)  # type: ignore[union-attr]


def _bookmarked_sentence(path: Path, text: str) -> None:
    document = PyDocxDocument()
    paragraph = document.add_paragraph()
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "0")
    start.set(qn("w:name"), "span")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "0")
    paragraph._p.append(start)
    _append_run(paragraph._p, text)
    paragraph._p.append(end)
    _write_docx(path, document)


def _hyperlinked_sentence(path: Path, text: str) -> None:
    document = PyDocxDocument()
    paragraph = document.add_paragraph()
    rel_id = paragraph.part.relate_to("https://example.com/span", RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    hyperlink.set(qn("w:history"), "1")
    _append_run(hyperlink, text)
    paragraph._p.append(hyperlink)
    _write_docx(path, document)


def _sentence_with_drawing_inside_world(path: Path) -> None:
    document = PyDocxDocument()
    paragraph = document.add_paragraph()
    _append_run(paragraph._p, "Hello wo")
    drawing_run = OxmlElement("w:r")
    drawing_run.append(OxmlElement("w:drawing"))
    paragraph._p.append(drawing_run)
    _append_run(paragraph._p, "rld foo")
    _write_docx(path, document)


def test_open_path_delete_revision_in_middle_of_comment_removes_world(tmp_path: Path) -> None:
    path = tmp_path / "comment-world.docx"
    path.write_bytes(_plain_docx("Hello world foo"))
    document = DocxDocument.open(path)
    document.add_comment(CommentRange("body:p:0", 0, 15, "Hello world foo"), "note", AUTHOR)
    deleted = document.delete_revision(RevisionRange("body:p:0", 6, 11, "world"), REVIEWER)
    assert deleted.receipt.status is OperationStatus.APPLIED
    document.publish()

    assert _del_text(path.read_bytes()) == ["world"]
    accepted = accept_all_revisions_bytes(path.read_bytes(), drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "Hello  foo"


def test_open_path_delete_revision_in_middle_of_comment_wraps_lo_wo(tmp_path: Path) -> None:
    path = tmp_path / "comment-lo-wo.docx"
    path.write_bytes(_plain_docx("Hello world"))
    document = DocxDocument.open(path)
    document.add_comment(CommentRange("body:p:0", 0, 11, "Hello world"), "note", AUTHOR)
    deleted = document.delete_revision(RevisionRange("body:p:0", 3, 8, "lo wo"), REVIEWER)
    assert deleted.receipt.status is OperationStatus.APPLIED
    document.publish()

    assert _del_text(path.read_bytes()) == ["lo wo"]
    accepted = accept_all_revisions_bytes(path.read_bytes(), drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "Helrld"


def test_open_path_insert_revision_in_middle_of_comment(tmp_path: Path) -> None:
    path = tmp_path / "comment-insert.docx"
    path.write_bytes(_plain_docx("Hello world foo"))
    document = DocxDocument.open(path)
    document.add_comment(CommentRange("body:p:0", 0, 15, "Hello world foo"), "note", AUTHOR)
    inserted = document.insert_revision(RevisionPosition("body:p:0", 8), "X", REVIEWER)
    assert inserted.receipt.status is OperationStatus.APPLIED
    document.publish()

    assert _ins_text(path.read_bytes()) == ["X"]
    accepted = accept_all_revisions_bytes(path.read_bytes(), drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "Hello woXrld foo"
    reopened = DocxDocument.open(path)
    assert reopened.comments[0].anchor_text == "Hello woXrld foo"


def test_open_path_replace_revision_in_middle_of_comment(tmp_path: Path) -> None:
    path = tmp_path / "comment-replace.docx"
    path.write_bytes(_plain_docx("Hello world foo"))
    document = DocxDocument.open(path)
    document.add_comment(CommentRange("body:p:0", 0, 15, "Hello world foo"), "note", AUTHOR)
    deleted, inserted = document.replace_revision(
        RevisionRange("body:p:0", 6, 11, "world"),
        "there",
        REVIEWER,
    )
    assert deleted.receipt.status is OperationStatus.APPLIED
    assert inserted.receipt.status is OperationStatus.APPLIED
    document.publish()

    assert _del_text(path.read_bytes()) == ["world"]
    assert _ins_text(path.read_bytes()) == ["there"]
    accepted = accept_all_revisions_bytes(path.read_bytes(), drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "Hello there foo"


def test_open_path_tracked_edits_in_middle_of_bookmark(tmp_path: Path) -> None:
    deleted_path = tmp_path / "bookmark-del.docx"
    _bookmarked_sentence(deleted_path, "Hello world foo")
    deleted_doc = DocxDocument.open(deleted_path)
    deleted = deleted_doc.delete_revision(RevisionRange("body:p:0", 6, 11, "world"), REVIEWER)
    assert deleted.receipt.status is OperationStatus.APPLIED
    deleted_doc.publish()
    assert _del_text(deleted_path.read_bytes()) == ["world"]
    accepted = accept_all_revisions_bytes(deleted_path.read_bytes(), drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "Hello  foo"

    edited_path = tmp_path / "bookmark-edit.docx"
    _bookmarked_sentence(edited_path, "Hello world foo")
    edited = DocxDocument.open(edited_path)
    edited.replace_revision(RevisionRange("body:p:0", 6, 11, "world"), "there", REVIEWER)
    edited.insert_revision(RevisionPosition("body:p:0", 3), "X", REVIEWER)
    edited.publish()
    payload = edited_path.read_bytes()
    assert "world" in _del_text(payload)
    assert "there" in _ins_text(payload)
    assert "X" in _ins_text(payload)
    accepted_edit = accept_all_revisions_bytes(payload, drop_comments=False)
    assert _visible_text(accepted_edit.output_bytes) == "HelXlo there foo"


def test_open_path_insert_and_replace_on_paragraph_with_existing_revision(tmp_path: Path) -> None:
    path = tmp_path / "marked-revision.docx"
    path.write_bytes(_plain_docx("Hello world foo"))
    document = DocxDocument.open(path)
    document.insert_revision(RevisionPosition("body:p:0", 15), "!", REVIEWER)
    document.delete_revision(RevisionRange("body:p:0", 6, 11, "world"), REVIEWER)
    document.replace_revision(RevisionRange("body:p:0", 0, 5, "Hello"), "Hey", REVIEWER)
    document.publish()

    payload = path.read_bytes()
    assert "world" in _del_text(payload)
    assert "Hello" in _del_text(payload)
    assert "Hey" in _ins_text(payload)
    accepted = accept_all_revisions_bytes(payload, drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "Hey  foo!"


def test_open_path_tracked_edits_in_middle_of_existing_insertion(tmp_path: Path) -> None:
    path = tmp_path / "inside-ins.docx"
    path.write_bytes(_plain_docx("Hello world foo"))
    document = DocxDocument.open(path)
    document.insert_revision(RevisionPosition("body:p:0", 0), "ABCDEF", REVIEWER)
    document.delete_revision(RevisionRange("body:p:0", 2, 4, "CD"), REVIEWER)
    document.replace_revision(RevisionRange("body:p:0", 4, 6, "EF"), "YZ", REVIEWER)
    document.publish()

    payload = path.read_bytes()
    assert "CD" in _del_text(payload)
    assert "EF" in _del_text(payload)
    assert "YZ" in _ins_text(payload)
    accepted = accept_all_revisions_bytes(payload, drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "ABYZHello world foo"


def test_open_path_tracked_edits_in_middle_of_hyperlink(tmp_path: Path) -> None:
    path = tmp_path / "hyperlink.docx"
    _hyperlinked_sentence(path, "Hello world foo")
    document = DocxDocument.open(path)
    document.insert_revision(RevisionPosition("body:p:0", 3), "X", REVIEWER)
    document.delete_revision(RevisionRange("body:p:0", 7, 12, "world"), REVIEWER)
    document.replace_revision(RevisionRange("body:p:0", 13, 16, "foo"), "bar", REVIEWER)
    document.publish()

    payload = path.read_bytes()
    root = _part_xml(payload)
    assert root.xpath("count(//w:hyperlink//w:del)", namespaces=NS) >= 1
    assert root.xpath("count(//w:hyperlink//w:ins)", namespaces=NS) >= 1
    assert "world" in _del_text(payload)
    assert "foo" in _del_text(payload)
    assert "bar" in _ins_text(payload)
    accepted = accept_all_revisions_bytes(payload, drop_comments=False)
    assert _visible_text(accepted.output_bytes) == "HelXlo  bar"


def test_open_path_refuses_revision_that_crosses_a_drawing(tmp_path: Path) -> None:
    path = tmp_path / "drawing.docx"
    _sentence_with_drawing_inside_world(path)
    original = path.read_bytes()
    document = DocxDocument.open(path)
    before = document.to_bytes()

    with pytest.raises(RevisionMutationError, match="opaque content"):
        document.delete_revision(RevisionRange("body:p:0", 6, 11, "world"), REVIEWER)

    assert document.to_bytes() == before
    assert path.read_bytes() == original
    with pytest.raises(RevisionMutationError, match="opaque content"):
        document.replace_revision(RevisionRange("body:p:0", 6, 11, "world"), "there", REVIEWER)
    assert document.to_bytes() == before
    assert path.read_bytes() == original
