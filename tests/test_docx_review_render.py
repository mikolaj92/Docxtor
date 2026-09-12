from __future__ import annotations

from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from docxtor import (
    PhysicalReviewEdit,
    PhysicalReviewPlan,
    render_physical_clean,
    render_physical_review,
)
from docxtor.docx_review_inventory import inventory_review_markup
from docxtor.docx_review_models import ReviewCoverage

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def _docx_bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _styled_source() -> tuple[bytes, str]:
    document = Document()
    style = document.styles.add_style("Synthetic Review Anchor", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Inches(0.25)
    run = paragraph.add_run("Anchor paragraph with explicit styling.")
    run.font.name = "Courier New"
    run.font.size = Pt(18)
    run.font.bold = True
    return _docx_bytes(document), style.style_id


def _insertion(text: str, *, comment_text: str | None = None) -> PhysicalReviewEdit:
    return PhysicalReviewEdit(
        action_id="inserted-paragraph",
        locator="body:p:0",
        operation="insert",
        start_offset=0,
        end_offset=0,
        replacement_text=text,
        comment_text=comment_text,
        new_paragraph=True,
    )


def _document_xml(path: Path) -> ElementTree.Element:
    with ZipFile(path) as archive:
        return ElementTree.fromstring(archive.read("word/document.xml"))


def _body_paragraphs(path: Path) -> list[ElementTree.Element]:
    body = _document_xml(path).find(f"{{{_W}}}body")
    assert body is not None
    return body.findall(f"{{{_W}}}p")


def _inserted_run_properties(paragraph: ElementTree.Element) -> ElementTree.Element | None:
    revision = paragraph.find(f"{{{_W}}}ins")
    assert revision is not None
    run = revision.find(f"{{{_W}}}r")
    assert run is not None
    return run.find(f"{{{_W}}}rPr")


def test_new_tracked_paragraph_inherits_anchor_paragraph_and_run_style(tmp_path: Path) -> None:
    source, style_id = _styled_source()
    output = tmp_path / "reviewed.docx"
    plan = PhysicalReviewPlan(
        edits=(
            _insertion(
                "Inserted text follows the anchor style.",
                comment_text="Review the inserted text.",
            ),
        )
    )

    render_physical_review(source, output, plan)

    anchor, inserted = _body_paragraphs(output)
    anchor_style = anchor.find(f"{{{_W}}}pPr/{{{_W}}}pStyle")
    assert anchor_style is not None
    assert anchor_style.get(f"{{{_W}}}val") == style_id
    ppr = inserted.find(f"{{{_W}}}pPr")
    assert ppr is not None
    inserted_style = ppr.find(f"{{{_W}}}pStyle")
    assert inserted_style is not None
    assert inserted_style.get(f"{{{_W}}}val") == style_id
    alignment = ppr.find(f"{{{_W}}}jc")
    assert alignment is not None
    assert alignment.get(f"{{{_W}}}val") == "center"
    indentation = ppr.find(f"{{{_W}}}ind")
    assert indentation is not None
    assert indentation.get(f"{{{_W}}}left") == "360"

    paragraph_mark = ppr.find(f"{{{_W}}}rPr/{{{_W}}}ins")
    assert paragraph_mark is not None
    assert not list(paragraph_mark)

    rpr = _inserted_run_properties(inserted)
    assert rpr is not None
    rfonts = rpr.find(f"{{{_W}}}rFonts")
    assert rfonts is not None
    assert rfonts.get(f"{{{_W}}}ascii") == "Courier New"
    assert rfonts.get(f"{{{_W}}}hAnsi") == "Courier New"
    font_size = rpr.find(f"{{{_W}}}sz")
    assert font_size is not None
    assert font_size.get(f"{{{_W}}}val") == "36"
    assert rpr.find(f"{{{_W}}}b") is not None

    inventory = inventory_review_markup(output.read_bytes())
    assert inventory.coverage is ReviewCoverage.COMPLETE
    assert any(revision.raw_kind == "ins" for revision in inventory.revisions)
    assert [comment.text for comment in inventory.comments] == ["Review the inserted text."]


@pytest.mark.parametrize(
    "run_history_only",
    [False, True],
    ids=["empty-run-properties", "revision-snapshot-only"],
)
def test_new_paragraph_uses_paragraph_mark_when_anchor_run_properties_are_unusable(
    tmp_path: Path, run_history_only: bool
) -> None:
    document = Document()
    paragraph = document.add_paragraph("Anchor text with no usable run overrides.")
    paragraph_mark_properties = OxmlElement("w:rPr")
    paragraph._p.get_or_add_pPr().append(paragraph_mark_properties)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Courier New")
    fonts.set(qn("w:hAnsi"), "Courier New")
    paragraph_mark_properties.append(fonts)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), "32")
    paragraph_mark_properties.append(font_size)

    run_properties = paragraph.runs[0]._r.get_or_add_rPr()
    if run_history_only:
        revision_snapshot = OxmlElement("w:rPrChange")
        revision_snapshot.set(qn("w:id"), "71")
        revision_snapshot.append(OxmlElement("w:rPr"))
        run_properties.append(revision_snapshot)

    source = _docx_bytes(document)
    output = tmp_path / "paragraph-mark-fallback.docx"
    render_physical_review(
        source,
        output,
        PhysicalReviewPlan(edits=(_insertion("Inserted text uses paragraph-mark formatting."),)),
    )

    inserted = _body_paragraphs(output)[1]
    rpr = _inserted_run_properties(inserted)
    assert rpr is not None
    rfonts = rpr.find(f"{{{_W}}}rFonts")
    assert rfonts is not None
    assert rfonts.get(f"{{{_W}}}ascii") == "Courier New"
    size = rpr.find(f"{{{_W}}}sz")
    assert size is not None
    assert size.get(f"{{{_W}}}val") == "32"


def test_new_paragraph_skips_anchor_run_with_direct_vanish(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    hidden = paragraph.add_run("Hidden anchor text.")
    hidden.font.name = "Times New Roman"
    hidden.font.size = Pt(12)
    hidden_properties = hidden._r.get_or_add_rPr()
    hidden_properties.append(OxmlElement("w:vanish"))

    visible = paragraph.add_run("Visible anchor text.")
    visible.font.name = "Courier New"
    visible.font.size = Pt(18)
    source = _docx_bytes(document)
    output = tmp_path / "direct-vanish.docx"
    render_physical_review(
        source,
        output,
        PhysicalReviewPlan(edits=(_insertion("Inserted text follows visible formatting."),)),
    )

    rpr = _inserted_run_properties(_body_paragraphs(output)[1])
    assert rpr is not None
    rfonts = rpr.find(f"{{{_W}}}rFonts")
    assert rfonts is not None
    assert rfonts.get(f"{{{_W}}}ascii") == "Courier New"
    size = rpr.find(f"{{{_W}}}sz")
    assert size is not None
    assert size.get(f"{{{_W}}}val") == "36"


def test_new_paragraph_keeps_anchor_run_with_disabled_direct_vanish(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("Visible because vanish is disabled.")
    first.font.name = "Times New Roman"
    first.font.size = Pt(12)
    vanish = OxmlElement("w:vanish")
    vanish.set(qn("w:val"), "0")
    first._r.get_or_add_rPr().append(vanish)

    second = paragraph.add_run("The later run is not the anchor style.")
    second.font.name = "Courier New"
    second.font.size = Pt(18)
    output = tmp_path / "disabled-direct-vanish.docx"
    render_physical_review(
        _docx_bytes(document),
        output,
        PhysicalReviewPlan(edits=(_insertion("Inserted text follows the first run."),)),
    )

    rpr = _inserted_run_properties(_body_paragraphs(output)[1])
    assert rpr is not None
    rfonts = rpr.find(f"{{{_W}}}rFonts")
    assert rfonts is not None
    assert rfonts.get(f"{{{_W}}}ascii") == "Times New Roman"
    size = rpr.find(f"{{{_W}}}sz")
    assert size is not None
    assert size.get(f"{{{_W}}}val") == "24"


def test_new_paragraph_without_explicit_anchor_style_uses_document_default(
    tmp_path: Path,
) -> None:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(14)
    anchor = document.add_paragraph("Anchor uses the implicit default paragraph style.")
    assert anchor._p.pPr is None
    source = _docx_bytes(document)
    output = tmp_path / "default-style.docx"

    render_physical_review(
        source,
        output,
        PhysicalReviewPlan(edits=(_insertion("Inserted paragraph uses the document default."),)),
    )

    _, inserted = _body_paragraphs(output)
    ppr = inserted.find(f"{{{_W}}}pPr")
    assert ppr is not None
    assert ppr.find(f"{{{_W}}}pStyle") is None
    assert _inserted_run_properties(inserted) is None

    with ZipFile(output) as archive:
        styles = ElementTree.fromstring(archive.read("word/styles.xml"))
    default_style = next(
        style
        for style in styles.findall(f"{{{_W}}}style")
        if style.get(f"{{{_W}}}styleId") == "Normal"
    )
    default_rpr = default_style.find(f"{{{_W}}}rPr")
    assert default_rpr is not None
    assert default_rpr.find(f"{{{_W}}}rFonts").get(f"{{{_W}}}ascii") == "Times New Roman"
    assert default_rpr.find(f"{{{_W}}}sz").get(f"{{{_W}}}val") == "28"


def test_new_paragraph_omits_direct_numbering_section_and_identity_markup(
    tmp_path: Path,
) -> None:
    source, _style_id = _styled_source()
    document = Document(BytesIO(source))
    anchor = document.paragraphs[0]
    anchor._p.set(f"{{{_W14}}}paraId", "AB12CD34")
    anchor._p.set(f"{{{_W14}}}textId", "AB12CD35")
    anchor._p.set(f"{{{_W}}}rsidR", "00000001")

    ppr = anchor._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "42")
    num_pr.extend((ilvl, num_id))
    style = ppr.find(qn("w:pStyle"))
    ppr.insert(ppr.index(style) + 1 if style is not None else 0, num_pr)

    section = OxmlElement("w:sectPr")
    page_size = OxmlElement("w:pgSz")
    page_size.set(qn("w:w"), "12240")
    page_size.set(qn("w:h"), "15840")
    section.append(page_size)
    ppr.append(section)

    ppr_change = OxmlElement("w:pPrChange")
    ppr_change.set(qn("w:id"), "91")
    ppr_change.append(OxmlElement("w:pPr"))
    ppr.append(ppr_change)

    run_properties = anchor.runs[0]._r.get_or_add_rPr()
    rpr_change = OxmlElement("w:rPrChange")
    rpr_change.set(qn("w:id"), "92")
    rpr_change.append(OxmlElement("w:rPr"))
    run_properties.append(rpr_change)
    source = _docx_bytes(document)
    output = tmp_path / "without-anchor-structure.docx"

    render_physical_review(
        source,
        output,
        PhysicalReviewPlan(edits=(_insertion("A new paragraph."),)),
    )

    inserted = _body_paragraphs(output)[1]
    ppr = inserted.find(f"{{{_W}}}pPr")
    assert ppr is not None
    assert ppr.find(f"{{{_W}}}sectPr") is None
    assert ppr.find(f"{{{_W}}}numPr") is None
    assert ppr.find(f"{{{_W}}}pPrChange") is None
    assert inserted.get(f"{{{_W14}}}paraId") is None
    assert inserted.get(f"{{{_W14}}}textId") is None
    assert inserted.get(f"{{{_W}}}rsidR") is None

    rpr = _inserted_run_properties(inserted)
    assert rpr is not None
    assert rpr.find(f"{{{_W}}}rPrChange") is None


def test_inline_review_comment_and_clean_render_remain_supported(tmp_path: Path) -> None:
    source, style_id = _styled_source()
    edit = PhysicalReviewEdit(
        action_id="inline-change",
        locator="body:p:0",
        operation="replace",
        start_offset=0,
        end_offset=6,
        replacement_text="Reviewed",
        expected_text="Anchor",
        comment_text="Check this inline change.",
    )
    reviewed = tmp_path / "inline-reviewed.docx"

    render_physical_review(source, reviewed, PhysicalReviewPlan(edits=(edit,)))

    inventory = inventory_review_markup(reviewed.read_bytes())
    assert inventory.coverage is ReviewCoverage.COMPLETE
    assert {revision.raw_kind for revision in inventory.revisions} == {"del", "ins"}
    assert [comment.text for comment in inventory.comments] == ["Check this inline change."]
    assert _body_paragraphs(reviewed)[0].find(f"{{{_W}}}pPr/{{{_W}}}pStyle").get(
        f"{{{_W}}}val"
    ) == style_id

    clean = tmp_path / "inline-clean.docx"
    clean_edit = PhysicalReviewEdit(
        action_id="inline-change",
        locator="body:p:0",
        operation="replace",
        start_offset=0,
        end_offset=6,
        replacement_text="Reviewed",
        expected_text="Anchor",
    )
    render_physical_clean(source, clean, (clean_edit,))

    assert Document(clean).paragraphs[0].text == "Reviewed paragraph with explicit styling."
    clean_inventory = inventory_review_markup(clean.read_bytes())
    assert clean_inventory.coverage is ReviewCoverage.COMPLETE
    assert clean_inventory.revisions == ()
    assert clean_inventory.comments == ()
