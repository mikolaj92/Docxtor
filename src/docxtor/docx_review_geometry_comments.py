"""Native point and multi-paragraph comment authoring from source-bound geometry."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO
from posixpath import basename, dirname, join, normpath
from typing import Any
from zipfile import ZipFile

from docx import Document as PyDocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from lxml import etree

from .docx_comment_mutations import CommentAuthor, CommentMutationError, CommentMutationResult
from .docx_ns import _TEXT_NODE_TAGS, CT_NS, REL_NS
from .docx_package import PackageEntry, parse_package_xml, read_package_entries
from .docx_review_geometry import (
    PhysicalCommentRange,
    _GeometryState,
    _marker_coordinates,
    _read_geometry_state,
)
from .docx_review_models import OperationReceipt, OperationStatus, ReviewCoverage
from .docx_stories import index_stories
from .docx_units import _set_text_node

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_COMMENTS_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"


def add_physical_comment_bytes(
    data: bytes,
    target: PhysicalCommentRange,
    text: str,
    author: CommentAuthor,
) -> CommentMutationResult:
    """Add one exact point or contiguous range comment without rewriting other parts.

    Every supplied slice must match the immutable input and provider paragraph order.
    Empty paragraphs are valid point/range boundaries. Interior opaque payloads stay
    intact; an endpoint inside an opaque wrapper cannot be represented and fails.
    """
    if not text:
        raise CommentMutationError("comment text must not be empty")
    state = _read_geometry_state(data)
    _validate_target(state, target)
    parts = {entry.name: entry for entry in state.entries}
    changed: dict[str, Any] = {}
    comments_name, comments_root = _comments_part(state, parts, changed)
    used_ids = {
        element.get(qn("w:id"))
        for root in state.roots.values()
        for element in root.iter()
        if element.tag
        in {
            qn("w:comment"),
            qn("w:commentRangeStart"),
            qn("w:commentRangeEnd"),
            qn("w:commentReference"),
        }
    }
    numeric_ids = [int(value) for value in used_ids if value is not None]
    candidate = max(numeric_ids, default=-1) + 1
    if candidate >= 2**31:
        candidate = 0
        while str(candidate) in used_ids:
            candidate += 1
    if candidate >= 2**31:
        raise CommentMutationError("no unused native comment identity is available")
    comment_id = str(candidate)
    _append_comment(comments_root, comment_id, text, author)
    changed[comments_name] = comments_root
    first, last = target.spans[0], target.spans[-1]
    start = state.paragraphs[first.locator]
    end = state.paragraphs[last.locator]
    start_marker = _marker("commentRangeStart", comment_id)
    end_marker = _marker("commentRangeEnd", comment_id)
    reference = etree.Element(qn("w:r"))
    reference.append(_marker("commentReference", comment_id))
    if first.locator == last.locator and first.start_offset == last.end_offset:
        _insert_boundary(start.element, first.start_offset, (start_marker, end_marker, reference))
    else:
        # Reverse endpoint order avoids invalidating an earlier offset when splitting one run.
        _insert_boundary(end.element, last.end_offset, (end_marker, reference))
        _insert_boundary(start.element, first.start_offset, (start_marker,))
    changed[start.geometry.part_name] = state.roots[start.geometry.part_name]
    changed[end.geometry.part_name] = state.roots[end.geometry.part_name]
    final_entries = tuple(
        replace(entry, data=_xml_bytes(changed[entry.name])) if entry.name in changed else entry
        for entry in state.entries
    ) + tuple(
        PackageEntry(name=name, data=_xml_bytes(root))
        for name, root in changed.items()
        if name not in parts
    )
    payload = _package_bytes(final_entries)
    after = _read_geometry_state(payload)
    created = next(
        (anchor for anchor in after.projection.comment_anchors if anchor.comment_id == comment_id),
        None,
    )
    if (
        after.projection.coverage is not ReviewCoverage.COMPLETE
        or created is None
        or not created.addressable
        or created.spans != target.spans
    ):
        raise CommentMutationError("physical comment creation was not confirmed after round-trip")
    for old in state.projection.comment_anchors:
        kept = next(
            (
                item
                for item in after.projection.comment_anchors
                if item.comment_id == old.comment_id
            ),
            None,
        )
        if kept is None or replace(kept, document_sha256=old.document_sha256) != old:
            raise CommentMutationError("existing physical comment anchors were not preserved")
    for old in state.projection.paragraphs:
        kept = after.paragraphs.get(old.locator)
        if kept is None or kept.geometry.text != old.text or kept.geometry.story_id != old.story_id:
            raise CommentMutationError("existing paragraph text or identity was not preserved")
    serialized_parts = {entry.name: entry.data for entry in read_package_entries(payload)}
    if any(
        serialized_parts.get(entry.name) != entry.data
        for entry in state.entries
        if entry.name not in changed
    ):
        raise CommentMutationError("untouched package entry bytes were not preserved")
    comments = tuple(index_stories(PyDocxDocument(BytesIO(payload))).comments)
    new_body = next((comment for comment in comments if comment.comment_id == comment_id), None)
    if (
        new_body is None
        or new_body.text != text
        or new_body.author != author.author
        or (author.initials is not None and new_body.initials != author.initials)
        or (author.date is not None and new_body.date != author.date)
    ):
        raise CommentMutationError("physical comment body was not confirmed after round-trip")
    affected = tuple(
        sorted(
            entry.name
            for entry in final_entries
            if entry.name not in parts or entry.data != parts[entry.name].data
        )
    )
    return CommentMutationResult(
        data=payload,
        receipt=OperationReceipt(
            operation="add_physical_comment",
            status=OperationStatus.APPLIED,
            affected_parts=affected,
            created_ids=(comment_id,),
            locator=first.locator,
            before_sha256=sha256(data).hexdigest(),
            after_sha256=sha256(payload).hexdigest(),
        ),
        comments=comments,
    )


def _validate_target(state: _GeometryState, target: PhysicalCommentRange) -> None:
    if state.projection.document_sha256 != target.document_sha256:
        raise CommentMutationError("physical comment geometry belongs to different source bytes")
    if state.projection.coverage is not ReviewCoverage.COMPLETE:
        raise CommentMutationError("physical comment source geometry is incomplete")
    if not target.spans:
        raise CommentMutationError("physical comment range must contain at least one span")
    records = []
    for span in target.spans:
        record = state.paragraphs.get(span.locator)
        if record is None:
            raise CommentMutationError(f"unknown physical comment locator: {span.locator}")
        if record.geometry.story_kind == "comment":
            raise CommentMutationError("comment bodies cannot host native comment anchors")
        value = record.geometry.text
        if (
            type(span.start_offset) is not int
            or type(span.end_offset) is not int
            or not 0 <= span.start_offset <= span.end_offset <= len(value)
            or value[span.start_offset : span.end_offset] != span.expected_text
        ):
            raise CommentMutationError(f"physical comment range changed at {span.locator}")
        records.append(record)
    first, last = records[0].geometry, records[-1].geometry
    if len({record.geometry.story_id for record in records}) != 1:
        raise CommentMutationError("physical comment range crosses stories")
    expected = tuple(
        record.geometry.locator
        for record in state.paragraphs.values()
        if record.geometry.story_id == first.story_id
        and first.story_order_index <= record.geometry.story_order_index <= last.story_order_index
    )
    if tuple(span.locator for span in target.spans) != expected:
        raise CommentMutationError(
            "physical comment range is reversed, duplicated or noncontiguous"
        )
    if len(records) > 1:
        for index, (span, record) in enumerate(zip(target.spans, records, strict=True)):
            if index and span.start_offset != 0:
                raise CommentMutationError("physical comment range omits an interior start")
            if index < len(records) - 1 and span.end_offset != len(record.geometry.text):
                raise CommentMutationError("physical comment range omits an interior end")


def _comments_part(
    state: _GeometryState,
    entries: dict[str, PackageEntry],
    changed: dict[str, Any],
) -> tuple[str, Any]:
    main = state.main_part_name
    rels_name = join(dirname(main), "_rels", basename(main) + ".rels")
    rels = (
        parse_package_xml(entries[rels_name].data, part_name=rels_name)
        if rels_name in entries
        else etree.Element(f"{{{REL_NS}}}Relationships", nsmap={None: REL_NS})
    )
    related = [node for node in rels if node.get("Type") == RT.COMMENTS]
    if len(related) > 1:
        raise CommentMutationError("multiple native comments relationships")
    if related:
        relationship = related[0]
        if relationship.get("TargetMode") == "External" or not relationship.get("Target"):
            raise CommentMutationError("native comments relationship is not an internal part")
        target = relationship.get("Target")
        assert target is not None
        name = normpath(join(dirname(main), target)) if not target.startswith("/") else target[1:]
        if name not in entries:
            raise CommentMutationError("native comments relationship target is missing")
        root = parse_package_xml(entries[name].data, part_name=name)
        if root.tag != qn("w:comments"):
            raise CommentMutationError("native comments relationship target has the wrong root")
        return name, root
    number = 1
    name = join(dirname(main), "comments.xml")
    while name in entries:
        number += 1
        name = join(dirname(main), f"comments{number}.xml")
    ids = {node.get("Id") for node in rels}
    rid = 1
    while f"rId{rid}" in ids:
        rid += 1
    etree.SubElement(
        rels,
        f"{{{REL_NS}}}Relationship",
        {
            "Id": f"rId{rid}",
            "Type": RT.COMMENTS,
            "Target": basename(name),
        },
    )
    changed[rels_name] = rels
    types_name = "[Content_Types].xml"
    if types_name not in entries:
        raise CommentMutationError("DOCX content types part is missing")
    types = parse_package_xml(entries[types_name].data, part_name=types_name)
    if any(node.get("PartName") == f"/{name}" for node in types):
        raise CommentMutationError("new comments part already has a content type declaration")
    etree.SubElement(
        types,
        f"{{{CT_NS}}}Override",
        {
            "PartName": f"/{name}",
            "ContentType": _COMMENTS_TYPE,
        },
    )
    changed[types_name] = types
    return name, etree.Element(qn("w:comments"), nsmap={"w": _W_NS})


def _append_comment(root: Any, comment_id: str, text: str, author: CommentAuthor) -> None:
    comment = etree.SubElement(
        root,
        qn("w:comment"),
        {
            qn("w:id"): comment_id,
            qn("w:author"): author.author,
            qn("w:date"): author.date or datetime.now(UTC).isoformat().replace("+00:00", "Z"),
        },
    )
    if author.initials is not None:
        comment.set(qn("w:initials"), author.initials)
    paragraph = etree.SubElement(comment, qn("w:p"))
    run = etree.SubElement(paragraph, qn("w:r"))
    for index, line in enumerate(text.split("\n")):
        if index:
            etree.SubElement(run, qn("w:br"))
        _set_text_node(etree.SubElement(run, qn("w:t")), line)


def _marker(kind: str, comment_id: str) -> Any:
    return etree.Element(qn(f"w:{kind}"), {qn("w:id"): comment_id})


def _node_width(element: Any) -> int:
    return len(_marker_coordinates(element)[1])


def _insert_boundary(paragraph: Any, offset: int, nodes: tuple[Any, ...]) -> None:
    cursor = 0
    for child in list(paragraph):
        if child.tag == qn("w:pPr"):
            continue
        width = _node_width(child)
        if offset == cursor:
            index = paragraph.index(child)
            for node in nodes:
                paragraph.insert(index, node)
                index += 1
            return
        if cursor < offset < cursor + width:
            if child.tag != qn("w:r"):
                raise CommentMutationError("physical comment endpoint lies inside opaque content")
            left, right = _split_run(child, offset - cursor)
            index = paragraph.index(child)
            paragraph.remove(child)
            for node in (left, *nodes, right):
                paragraph.insert(index, node)
                index += 1
            return
        cursor += width
    if offset != cursor:
        raise CommentMutationError("physical comment endpoint cannot be mapped to source XML")
    paragraph.extend(nodes)


def _split_run(run: Any, offset: int) -> tuple[Any, Any]:
    left, right = deepcopy(run), deepcopy(run)
    for copy in (left, right):
        for child in list(copy):
            if child.tag != qn("w:rPr"):
                copy.remove(child)
    cursor = 0
    for child in run:
        if child.tag == qn("w:rPr"):
            continue
        width = _node_width(child)
        if cursor < offset < cursor + width:
            if child.tag not in _TEXT_NODE_TAGS:
                raise CommentMutationError(
                    "physical comment endpoint lies inside opaque run content"
                )
            cut = offset - cursor
            before, after = deepcopy(child), deepcopy(child)
            _set_text_node(before, (child.text or "")[:cut])
            _set_text_node(after, (child.text or "")[cut:])
            left.append(before)
            right.append(after)
        elif cursor + width <= offset:
            left.append(deepcopy(child))
        else:
            right.append(deepcopy(child))
        cursor += width
    if not 0 < offset < cursor:
        raise CommentMutationError("physical comment endpoint cannot split this run")
    return left, right


def _xml_bytes(root: Any) -> bytes:
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def _package_bytes(entries: tuple[PackageEntry, ...]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for entry in entries:
            archive.writestr(entry.zip_info(), entry.data)
    payload = buffer.getvalue()
    read_package_entries(payload)
    return payload
