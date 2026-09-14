"""One-side mechanical projection used by the read-only DOCX comparator."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path

from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.styles.styles import Styles
from docx.text.paragraph import Paragraph
from lxml import etree

from .common import DocumentError
from .docx import DocxDocument
from .docx_compare_models import (
    CommentAnchor,
    ComparisonCoverage,
    ComparisonDiagnostic,
    DocumentBlock,
    DocumentComment,
    DocumentSpan,
    DocumentView,
    FormattingSpan,
    RevisionReference,
)
from .docx_facts import DocxFactsSnapshot, ParagraphFact, docx_facts
from .docx_models import AddressableSpan
from .docx_review_inventory import inventory_review_markup
from .docx_review_models import ReviewMarkupInventory
from .docx_revisions import Revision


@dataclass(frozen=True)
class ProjectedDocument:
    view: DocumentView
    revisions: tuple[RevisionReference, ...]
    style_parts: tuple[tuple[str, str], ...]


_UNSUPPORTED_INLINE_LOCAL_NAMES = {
    "endnoteReference",
    "fldChar",
    "fldSimple",
    "footnoteReference",
    "noBreakHyphen",
    "oMath",
    "oMathPara",
    "softHyphen",
    "sym",
}


def project_document(source: str | Path | bytes, *, side: str) -> ProjectedDocument:
    """Read a document through existing Docxtor facts and review inventory APIs."""
    data = _read_source(source)
    document = DocxDocument.open_bytes(data)
    facts = docx_facts(data)
    review = inventory_review_markup(data)
    sha = sha256(data).hexdigest()
    diagnostics = _coverage_diagnostics(facts, review, side)
    spans_by_container = _spans_by_container(document)
    styles = document._doc.styles
    revision_kinds_by_id: dict[tuple[str, str], set[str]] = {}
    revision_counts_by_id: dict[tuple[str, str], int] = {}
    for revision in review.revisions:
        if revision.revision_id is None:
            continue
        key = (revision.part_name, revision.revision_id)
        revision_kinds_by_id.setdefault(key, set()).add(revision.raw_kind)
        revision_counts_by_id[key] = revision_counts_by_id.get(key, 0) + 1
    for (part_name, revision_id), count in sorted(revision_counts_by_id.items()):
        if count > 1:
            diagnostics.append(
                ComparisonDiagnostic(
                    "duplicate_revision_identity",
                    f"{count} revision wrappers share id {revision_id}; "
                    "exact identity is ambiguous",
                    side,
                    part_name,
                )
            )
    blocks: list[DocumentBlock] = []

    for fact in facts.paragraphs:
        if fact.container_id.startswith("comment:"):
            continue
        paragraph = document.resolve_paragraph(fact.container_id)
        if paragraph is None:
            diagnostics.append(
                ComparisonDiagnostic(
                    "paragraph_projection_unavailable",
                    "paragraph facts have no addressable document block",
                    side,
                    fact.part_name,
                    fact.container_id,
                )
            )
            continue
        diagnostics.extend(_unsupported_inline_diagnostics(paragraph, fact, side))
        block_spans = []
        accepted_parts: list[str] = []
        accepted_offset = 0
        raw_parts: list[str] = []
        for span in sorted(
            spans_by_container.get(fact.container_id, ()), key=lambda x: x.start_offset
        ):
            raw_parts.append(span.text)
            revision_kinds = revision_kinds_by_id.get(
                (fact.part_name or "", span.revision_id or ""), set()
            )
            revision_kind = next(iter(revision_kinds)) if len(revision_kinds) == 1 else None
            if revision_kind is None:
                if span.role == "insertion":
                    revision_kind = "ins"
                elif span.role == "deletion":
                    revision_kind = "del"
            include_in_accepted_text = revision_kind not in {"del", "moveFrom"}
            accepted_start = accepted_offset if include_in_accepted_text else None
            accepted_end = accepted_offset + len(span.text) if include_in_accepted_text else None
            if include_in_accepted_text:
                accepted_parts.append(span.text)
                accepted_offset += len(span.text)
            block_spans.append(
                DocumentSpan(
                    role=span.role,
                    text=span.text,
                    start_offset=span.start_offset,
                    end_offset=span.end_offset,
                    accepted_start_offset=accepted_start,
                    accepted_end_offset=accepted_end,
                    revision_kind=revision_kind,
                    revision_id=span.revision_id,
                    revision_author=span.revision_author,
                    revision_date=span.revision_date,
                    hyperlink_anchor=span.hyperlink_anchor,
                    hyperlink_rel_id=span.hyperlink_rel_id,
                )
            )
        raw_text = "".join(raw_parts)
        accepted_text = "".join(accepted_parts)
        if _has_numbering(paragraph, styles):
            diagnostics.append(
                ComparisonDiagnostic(
                    "numbering_not_projected",
                    "paragraph numbering is retained as a fact but is not rendered "
                    "in the text projection",
                    side,
                    fact.part_name,
                    fact.container_id,
                )
            )
        block_id = _stable_id("block", sha, fact.container_id)
        paragraph_formatting_sha, formatting_spans, formatting_sha = _formatting_projection(
            paragraph, fact
        )
        blocks.append(
            DocumentBlock(
                id=block_id,
                pair_id=_stable_id("pair", sha, block_id),
                order=len(blocks),
                story_id=_story_id(fact),
                coordinate=fact.coordinate,
                style_id=fact.style_id,
                outline_level=fact.outline_level,
                raw_text=raw_text,
                text=accepted_text,
                spans=tuple(block_spans),
                formatting_sha256=formatting_sha,
                paragraph_formatting_sha256=paragraph_formatting_sha,
                formatting_spans=formatting_spans,
            )
        )

    block_by_locator = {block.coordinate.container_id: block for block in blocks}
    comments: list[DocumentComment] = []
    for comment in document.comments:
        host_block = block_by_locator.get(comment.locator or "")
        anchor_start = anchor_end = None
        if comment.anchor_text and host_block is not None:
            # AddressableComment.anchor_text is the accepted visible selection;
            # keep its offsets in the same coordinate space as DocumentBlock.text.
            positions = _find_all(host_block.text, comment.anchor_text)
            if len(positions) == 1:
                anchor_start = positions[0]
                anchor_end = anchor_start + len(comment.anchor_text)
            else:
                diagnostics.append(
                    ComparisonDiagnostic(
                        "comment_text_anchor_unresolved",
                        "comment text anchor is absent or ambiguous in its paragraph",
                        side,
                        "word/comments.xml",
                        comment.locator,
                    )
                )
        elif host_block is None:
            diagnostics.append(
                ComparisonDiagnostic(
                    "comment_block_anchor_unresolved",
                    "comment has no addressable host paragraph in the projection",
                    side,
                    "word/comments.xml",
                    comment.locator,
                )
            )
        comments.append(
            DocumentComment(
                comment_id=comment.comment_id,
                text=comment.text,
                author=comment.author,
                initials=comment.initials,
                date=comment.date,
                parent_id=comment.parent_id,
                anchor=(
                    CommentAnchor(
                        comment.locator,
                        host_block.id if host_block is not None else None,
                        host_block.pair_id if host_block is not None else None,
                        comment.anchor_text,
                        anchor_start,
                        anchor_end,
                    )
                    if comment.locator or comment.anchor_text
                    else None
                ),
            )
        )

    revision_refs = _revision_references(
        facts,
        review,
        blocks,
        spans_by_container,
        side,
        diagnostics,
    )
    coverage = ComparisonCoverage.INCOMPLETE if diagnostics else ComparisonCoverage.COMPLETE
    style_parts = tuple(
        sorted(
            (part.name, part.sha256)
            for part in facts.parts
            if part.name.startswith("word/styles") or part.name.startswith("word/theme/")
        )
    )
    return ProjectedDocument(
        DocumentView(sha, tuple(blocks), tuple(comments), coverage, tuple(diagnostics)),
        tuple(revision_refs),
        style_parts,
    )


def _read_source(source: str | Path | bytes) -> bytes:
    if isinstance(source, bytes):
        return source
    try:
        return Path(source).read_bytes()
    except OSError as exc:
        raise DocumentError(f"cannot read DOCX source: {exc}") from exc


def _spans_by_container(document: DocxDocument) -> dict[str, list[AddressableSpan]]:
    result: dict[str, list[AddressableSpan]] = {}
    for span in document.spans:
        result.setdefault(span.container_id, []).append(span)
    return result


def _coverage_diagnostics(
    facts: DocxFactsSnapshot, review: ReviewMarkupInventory, side: str
) -> list[ComparisonDiagnostic]:
    diagnostics = [
        ComparisonDiagnostic(item.code, item.message, side, item.part_name)
        for item in facts.diagnostics
    ]
    diagnostics.extend(
        ComparisonDiagnostic(item.code, item.message, side, item.part_name)
        for item in review.diagnostics
    )
    if facts.coverage.value == "incomplete" and not facts.diagnostics:
        diagnostics.append(
            ComparisonDiagnostic(
                "facts_coverage_incomplete", "DOCX facts coverage is incomplete", side
            )
        )
    if review.coverage.value == "incomplete" and not review.diagnostics:
        diagnostics.append(
            ComparisonDiagnostic(
                "review_markup_coverage_incomplete",
                "revision and comment coverage is incomplete",
                side,
            )
        )
    for revision in review.revisions:
        if revision.raw_kind not in {"ins", "del", "rPrChange", "pPrChange"}:
            diagnostics.append(
                ComparisonDiagnostic(
                    "unsupported_revision_kind",
                    f"revision kind {revision.raw_kind} is inventoried but not "
                    "projected as accepted text",
                    side,
                    revision.part_name,
                    revision.locator,
                )
            )
    for item in (*facts.media, *facts.embedded_objects):
        diagnostics.append(
            ComparisonDiagnostic(
                "unsupported_object",
                "embedded or media object is inventoried but not rendered in the text projection",
                side,
                item.part_name,
                item.fact_id,
            )
        )
    return diagnostics


def _story_id(fact: ParagraphFact) -> str:
    coordinate = fact.coordinate
    if coordinate.story_kind == "table":
        return f"table:{coordinate.table_index or 0}"
    if coordinate.story_kind in {"body", "txbx"}:
        if coordinate.story_kind == "txbx":
            return fact.container_id.rsplit(":p:", 1)[0]
        return "body"
    return fact.container_id.rsplit(":p:", 1)[0]


def _has_numbering(paragraph: Paragraph, styles: Styles) -> bool:
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        return True
    style = styles.get_by_id(paragraph._p.style, WD_STYLE_TYPE.PARAGRAPH)
    seen_style_ids = set()
    while style is not None and style.style_id not in seen_style_ids:
        seen_style_ids.add(style.style_id)
        style_ppr = style._element.find(qn("w:pPr"))
        if style_ppr is not None and style_ppr.find(qn("w:numPr")) is not None:
            return True
        style = style.base_style
    return False


def _unsupported_inline_diagnostics(
    paragraph: Paragraph, fact: ParagraphFact, side: str
) -> list[ComparisonDiagnostic]:
    result = []
    tree = paragraph._p.getroottree()
    for element in paragraph._p.iter():
        if not isinstance(element.tag, str):
            continue
        local_name = etree.QName(element).localname
        if local_name not in _UNSUPPORTED_INLINE_LOCAL_NAMES:
            continue
        result.append(
            ComparisonDiagnostic(
                "unsupported_inline_surface",
                f"w:{local_name} can affect visible content but is not projected",
                side,
                fact.part_name,
                tree.getpath(element),
            )
        )
    return result


def _formatting_projection(
    paragraph: Paragraph, fact: ParagraphFact
) -> tuple[str, tuple[FormattingSpan, ...], str]:
    ppr = paragraph._p.find(qn("w:pPr"))
    ppr_bytes = _effective_properties(ppr)
    paragraph_signature = sha256(
        repr((fact.style_id, fact.outline_level, ppr_bytes.hex())).encode("utf-8")
    ).hexdigest()
    formatting_spans: list[FormattingSpan] = []
    offset = 0
    for run in paragraph._p.iter(qn("w:r")):
        if _inside_deleted_revision(run):
            continue
        text = _run_accepted_text(run)
        if not text:
            continue
        rpr = run.find(qn("w:rPr"))
        style_hash = sha256(_effective_properties(rpr)).hexdigest()
        if formatting_spans and formatting_spans[-1].properties_sha256 == style_hash:
            previous = formatting_spans[-1]
            formatting_spans[-1] = FormattingSpan(
                previous.start_offset, offset + len(text), style_hash
            )
        else:
            formatting_spans.append(FormattingSpan(offset, offset + len(text), style_hash))
        offset += len(text)
    formatted = tuple(formatting_spans)
    payload = repr(
        (
            paragraph_signature,
            tuple(
                (item.start_offset, item.end_offset, item.properties_sha256) for item in formatted
            ),
        )
    )
    return paragraph_signature, formatted, sha256(payload.encode("utf-8")).hexdigest()


def _effective_properties(element: etree._Element | None) -> bytes:
    if element is None:
        return b""
    clone = deepcopy(element)
    changed = {qn("w:rPrChange"), qn("w:pPrChange")}
    for candidate in tuple(clone.iter()):
        for child in tuple(candidate):
            if child.tag in changed:
                candidate.remove(child)
    return etree.tostring(clone, method="c14n")


def _inside_deleted_revision(element: etree._Element) -> bool:
    return any(
        ancestor.tag in {qn("w:del"), qn("w:moveFrom")} for ancestor in element.iterancestors()
    )


def _run_accepted_text(run: etree._Element) -> str:
    parts = []
    for node in run.iter():
        if _inside_deleted_revision(node):
            continue
        if node.tag in {qn("w:t"), qn("w:delText")}:
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        elif node.tag in {qn("w:br"), qn("w:cr")}:
            parts.append("\n")
    return "".join(parts)


def _revision_references(
    facts: DocxFactsSnapshot,
    review: ReviewMarkupInventory,
    blocks: list[DocumentBlock],
    spans_by_container: dict[str, list[AddressableSpan]],
    side: str,
    diagnostics: list[ComparisonDiagnostic],
) -> list[RevisionReference]:
    block_by_locator = {block.coordinate.container_id: block for block in blocks}
    scoped_counts: dict[tuple[str, str, str, str], int] = {}
    facts_by_revision: dict[tuple[str, str], ParagraphFact | None] = {}
    for revision in review.revisions:
        fact = _fact_for_revision_locator(facts.paragraphs, revision)
        facts_by_revision[(revision.part_name, revision.locator)] = fact
        if revision.revision_id is None or fact is None:
            continue
        key = (revision.part_name, revision.revision_id, revision.raw_kind, fact.container_id)
        scoped_counts[key] = scoped_counts.get(key, 0) + 1

    result = []
    for revision in review.revisions:
        fact = facts_by_revision.get((revision.part_name, revision.locator))
        candidate_facts = (
            (fact,)
            if fact is not None
            else tuple(item for item in facts.paragraphs if item.part_name == revision.part_name)
        )
        matching_spans = [
            (candidate, span)
            for candidate in candidate_facts
            for span in spans_by_container.get(candidate.container_id, ())
            if revision.revision_id is not None
            and span.revision_id == revision.revision_id
            and _span_matches_revision_kind(span, revision.raw_kind)
        ]
        unique_matching_facts = {item.container_id: item for item, _ in matching_spans}
        if fact is None and len(unique_matching_facts) == 1:
            fact = next(iter(unique_matching_facts.values()))
        if fact is None and len(unique_matching_facts) > 1:
            diagnostics.append(
                ComparisonDiagnostic(
                    "revision_block_anchor_ambiguous",
                    "revision id maps to spans in multiple paragraphs and has no "
                    "resolvable XML paragraph",
                    side,
                    revision.part_name,
                    revision.locator,
                )
            )
            matching_spans = []

        duplicate_scope = (
            (
                revision.part_name,
                revision.revision_id,
                revision.raw_kind,
                fact.container_id,
            )
            if revision.revision_id is not None and fact is not None
            else None
        )
        if duplicate_scope is not None and scoped_counts.get(duplicate_scope, 0) > 1:
            diagnostics.append(
                ComparisonDiagnostic(
                    "duplicate_revision_anchor",
                    "multiple revision wrappers share an id, kind, and paragraph; "
                    "text anchoring is ambiguous",
                    side,
                    revision.part_name,
                    revision.locator,
                )
            )
            matching_spans = []

        block = block_by_locator.get(fact.container_id) if fact is not None else None
        text_parts: list[str] = []
        offsets: list[tuple[int, int]] = []
        for _fact, span in sorted(
            matching_spans,
            key=lambda item: (item[1].start_offset, item[1].end_offset),
        ):
            text_parts.append(span.text)
            offsets.append((span.start_offset, span.end_offset))
        if block is None:
            diagnostics.append(
                ComparisonDiagnostic(
                    "revision_block_anchor_unresolved",
                    "revision is inventoried but has no addressable projected paragraph",
                    side,
                    revision.part_name,
                    revision.locator,
                )
            )
        result.append(
            RevisionReference(
                part_name=revision.part_name,
                locator=revision.locator,
                kind=revision.raw_kind,
                revision_id=revision.revision_id,
                author=revision.author,
                date=revision.date,
                container_id=block.coordinate.container_id if block is not None else None,
                block_id=block.id if block is not None else None,
                pair_id=block.pair_id if block is not None else None,
                text="".join(text_parts),
                start_offset=min((item[0] for item in offsets), default=None),
                end_offset=max((item[1] for item in offsets), default=None),
            )
        )
    return result


def _span_matches_revision_kind(span: AddressableSpan, kind: str) -> bool:
    expected_role = {"ins": "insertion", "del": "deletion"}.get(kind)
    return expected_role is not None and span.role == expected_role


def _fact_for_revision_locator(
    paragraphs: tuple[ParagraphFact, ...], revision: Revision
) -> ParagraphFact | None:
    candidates = [
        fact
        for fact in paragraphs
        if fact.part_name == revision.part_name
        and fact.xml_path is not None
        and revision.locator.startswith(fact.xml_path + "/")
    ]
    return max(candidates, key=lambda item: len(item.xml_path or ""), default=None)


def _find_all(text: str, target: str) -> list[int]:
    positions: list[int] = []
    start = 0
    while True:
        offset = text.find(target, start)
        if offset < 0:
            return positions
        positions.append(offset)
        start = offset + 1


def _stable_id(prefix: str, *parts: str) -> str:
    return f"{prefix}-{sha256(chr(31).join(parts).encode('utf-8')).hexdigest()[:20]}"
