"""Typed, domain-blind projection for DOCX review consumers."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from docx.oxml.ns import qn

from .docx import DocxDocument
from .docx_facts import ParagraphFact, docx_facts
from .docx_models import AddressableComment, AddressableSpan
from .docx_review_inventory import inventory_review_markup
from .docx_review_models import ReviewCoverage, ReviewDiagnostic

if TYPE_CHECKING:
    from .docx_review_geometry import PhysicalReviewGeometry


@dataclass(frozen=True)
class ReviewParagraphProjection:
    locator: str
    text: str
    paragraph_index: int | None
    story_kind: str
    is_heading: bool
    opaque_ranges: tuple[tuple[int, int], ...] = ()


@dataclass(frozen=True)
class ReviewNoteProjection:
    note_id: str
    kind: str
    text: str


@dataclass(frozen=True)
class DocxReviewProjection:
    paragraphs: tuple[ReviewParagraphProjection, ...]
    spans: tuple[AddressableSpan, ...]
    comments: tuple[AddressableComment, ...]
    notes: tuple[ReviewNoteProjection, ...]
    table_count: int
    coverage: ReviewCoverage
    diagnostics: tuple[ReviewDiagnostic, ...]
    paragraph_mark_revisions: tuple[AddressableSpan, ...] = ()
    physical_geometry: PhysicalReviewGeometry | None = None
    tracked_revisions_detected: bool = False


def project_docx_for_review(source: str | Path | bytes) -> DocxReviewProjection:
    from .docx_review_geometry import project_docx_review_geometry

    data = source if isinstance(source, bytes) else Path(source).read_bytes()
    geometry = project_docx_review_geometry(data)
    document = DocxDocument.open_bytes(data)
    inventory = inventory_review_markup(data)
    facts = docx_facts(data)
    paragraphs = []
    for paragraph in geometry.paragraphs:
        locator = paragraph.locator
        if not locator or locator.startswith(("comment:", "footnote:", "endnote:")):
            continue
        paragraphs.append(
            ReviewParagraphProjection(
                locator,
                paragraph.text,
                paragraph.paragraph_index,
                paragraph.story_kind,
                paragraph.is_heading,
                paragraph.opaque_ranges,
            )
        )
    notes = tuple(
        ReviewNoteProjection(
            item.value or item.fact_id, item.kind.split("_", 1)[0], item.target or ""
        )
        for item in facts.notes
        if item.kind in {"footnote_user", "endnote_user"}
    )
    table_ids = {
        fact.coordinate.table_index
        for fact in facts.paragraphs
        if fact.coordinate.table_index is not None
    }
    marks = _paragraph_mark_revisions(document, facts.paragraphs)
    coverage = inventory.coverage
    if geometry.coverage is ReviewCoverage.INCOMPLETE:
        coverage = ReviewCoverage.INCOMPLETE
    diagnostics = inventory.diagnostics + tuple(
        diagnostic for diagnostic in geometry.diagnostics if diagnostic not in inventory.diagnostics
    )
    effective_parts: dict[str, list[str]] = {}
    for span in document.spans:
        if span.role != "deletion":
            effective_parts.setdefault(span.container_id, []).append(span.text)
    # A deleted nonempty boundary can merge paragraphs and change their final
    # style. Physical receipts alone cannot map semantic actions back across
    # that merge. Retain fail-closed coverage until origin grouping is exposed.
    if any(
        mark.role == "deletion" and "".join(effective_parts.get(mark.container_id, ())).strip()
        for mark in marks
    ):
        coverage = ReviewCoverage.INCOMPLETE
        diagnostics += (
            ReviewDiagnostic(
                "unprojected_paragraph_merge",
                "A deleted nonempty paragraph boundary has no semantic origin mapping.",
            ),
        )
    if len(marks) != sum(
        revision.paragraph_mark and revision.raw_kind in {"ins", "del"}
        for revision in inventory.revisions
    ):
        coverage = ReviewCoverage.INCOMPLETE
        diagnostics += (
            ReviewDiagnostic(
                "unprojected_paragraph_mark_revision",
                "A paragraph-mark revision has no addressable review projection.",
            ),
        )
    return DocxReviewProjection(
        tuple(paragraphs),
        document.spans,
        inventory.comments,
        notes,
        len(table_ids),
        coverage,
        diagnostics,
        marks,
        geometry,
        bool(inventory.revisions)
        or any(
            diagnostic.code in {"unsupported_revision", "unsupported_namespace"}
            for diagnostic in inventory.diagnostics
        ),
    )


def _paragraph_mark_revisions(
    document: DocxDocument,
    paragraphs: tuple[ParagraphFact, ...],
) -> tuple[AddressableSpan, ...]:
    """Expose structural boundary revisions separately from textual spans."""
    marks = []
    offsets: dict[str, int] = {}
    for span in document.spans:
        offsets[span.container_id] = max(offsets.get(span.container_id, 0), span.end_offset)
    for segment in paragraphs:
        locator = segment.container_id
        paragraph = document.resolve_paragraph(locator)
        if paragraph is None:
            continue
        properties = paragraph._p.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
        if properties is None:
            continue
        offset = offsets.get(locator, 0)
        for index, node in enumerate(properties):
            if node.tag not in {qn("w:ins"), qn("w:del")}:
                continue
            marks.append(
                AddressableSpan(
                    span_id=f"{locator}:paragraph-mark:{index}",
                    container_id=locator,
                    role="insertion" if node.tag == qn("w:ins") else "deletion",
                    text="",
                    start_offset=offset,
                    end_offset=offset,
                    paragraph_index=segment.paragraph_index,
                    revision_id=node.get(qn("w:id")),
                    revision_author=node.get(qn("w:author")),
                    revision_date=node.get(qn("w:date")),
                )
            )
    return tuple(marks)
